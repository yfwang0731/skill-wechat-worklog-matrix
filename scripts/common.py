# -*- coding: utf-8 -*-
"""微信→需求跟踪矩阵 脚本共用模块：配置加载、路径自动识别、日期/列工具。

设计原则：
1. 不写死任何具体项目/人员/客户 —— 一切从 config.json 读取。
2. 本机路径一律自动识别（或由探测结果写入 config），不写死绝对路径。
3. 列位置靠"表头名"识别（probe.py 生成 column_mapping），不写死 A~AA。
"""
import os
import re
import sys
import json
import hashlib
import subprocess
from datetime import date

SKILL_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


# ---------------- 控制台编码 ----------------
def ensure_utf8_stdio():
    """把 stdout / stderr 切到 UTF-8，并尽量把 Windows 控制台也切过去。

    本 skill 的输出**全是中文**，而 Windows 上 Python 的标准流编码跟随控制台代码页
    （实测 GitHub 的 `windows-latest` runner 是 **cp1252**）：一 `print` 中文就
    `UnicodeEncodeError: 'charmap' codec can't encode ...`，**第一行输出就崩**。
    ubuntu / git-bash 都是 UTF-8，所以这个坑只在 Windows 上炸，本地极易漏。
    编码问题不该让工具崩掉 —— 显式切到 UTF-8 + `errors="replace"`。

    每个 `__main__` 入口都要在**任何输出之前**调用它（含 `smoke_test.py`）。
    CI 里**故意不设** `PYTHONUTF8` / `PYTHONIOENCODING`，否则就把这个缺陷盖住了。
    """
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except Exception:      # 流被替换过 / 不支持 reconfigure → 忽略，不影响主流程
            pass
    if os.name == "nt":
        # 让 cmd.exe / PowerShell 也按 UTF-8 解释这些字节，否则中文显示为乱码（等价于 chcp 65001）
        try:
            import ctypes
            ctypes.windll.kernel32.SetConsoleOutputCP(65001)
        except Exception:      # 无内核接口 / 非交互环境（CI 管道）
            pass


# ---------------- 配置 ----------------
def load_config(path=None):
    """加载 config.json。查找顺序：显式路径 → 环境变量 WM_CONFIG → 当前目录 → skill 根目录。"""
    cands = [path,
             os.environ.get("WM_CONFIG"),
             os.path.join(os.getcwd(), "config.json"),
             os.path.join(SKILL_ROOT, "config.json")]
    for p in cands:
        if p and os.path.isfile(p):
            # 走统一守卫：配置文件写坏（漏逗号之类）时给可执行提示，而不是裸 JSONDecodeError
            return load_json_file(p, "config 文件"), p
    return None, None


def col_index(letter):
    """列字母 -> 0-based 索引（A->0, AA->26）"""
    n = 0
    for ch in letter:
        n = n * 26 + (ord(ch.upper()) - 64)
    return n - 1


def col_letter(idx0):
    """0-based 索引 -> 列字母"""
    s, n = "", idx0 + 1
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


# ---------------- 路径自动识别 ----------------
def find_db_storage(preferred=None, hint=None):
    """自动定位微信 db_storage。多账号时可用 hint（账号目录特征串）优选。"""
    base = os.path.join(os.path.expanduser("~"), "Documents", "xwechat_files")
    if preferred:
        if os.path.isdir(preferred):
            return preferred
        cand = os.path.join(base, preferred, "db_storage")
        if os.path.isdir(cand):
            return cand
    if os.path.isdir(base):
        subs = [d for d in os.listdir(base)
                if os.path.isdir(os.path.join(base, d, "db_storage"))]
        if hint:
            for s in subs:
                if hint in s:
                    return os.path.join(base, s, "db_storage")
        if subs:
            return os.path.join(base, subs[0], "db_storage")
    raise FileNotFoundError(
        f"未在 {base} 找到微信 db_storage；请通过 --db-storage 指定或先跑 probe.py accounts。")


def find_wcdb_tool(explicit=None):
    """定位 wcdb-key-tool 的 windows 主脚本（第三方工具，不在 skill 内）。

    顺序：显式路径 → 环境变量 WCDB_KEY_TOOL → skill 内已知位置 → 家目录深度≤3 搜索。
    """
    if explicit and os.path.isfile(explicit):
        return explicit
    env = os.environ.get("WCDB_KEY_TOOL")
    if env and os.path.isfile(env):
        return env
    here = os.path.dirname(os.path.abspath(__file__))
    for c in [os.path.join(here, "tools", "wcdb-key-tool", "wcdb_key_tool_windows.py"),
              os.path.join(here, "wcdb-key-tool-main", "wcdb_key_tool_windows.py")]:
        if os.path.isfile(c):
            return os.path.abspath(c)
    home, target = os.path.expanduser("~"), "wcdb_key_tool_windows.py"
    for root, dirs, files in os.walk(home):
        if root[len(home):].count(os.sep) > 3:
            dirs[:] = []
            continue
        if target in files:
            return os.path.join(root, target)
    return None


# ---------------- 工具 ----------------
def md5hex(s):
    return hashlib.md5(s.encode("utf-8")).hexdigest()


def serial(dstr):
    """日期 -> Excel 序列号（1899-12-30 起）。容错支持 YYYY-MM-DD / YYYY/MM/DD /
    YYYY年M月D日 / 带时间（取前 10 位）等；无法解析返回 ''。"""
    if not dstr:
        return ""
    s = str(dstr).strip()
    m = re.match(r"^(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})", s)
    if not m:
        return ""
    try:
        y, mo, dd = map(int, m.groups())
        return (date(y, mo, dd) - date(1899, 12, 30)).days
    except Exception:
        return ""


def pick_sheet(workbook, sheet_hint=None):
    """按名字含 sheet_hint 选子表，找不到退回第一张。workbook 可为 openpyxl 或 GridWorkbook。"""
    names = list(workbook.sheetnames)
    hint = sheet_hint or ""
    cand = [n for n in names if hint and hint in n]
    return workbook[cand[0] if cand else names[0]]


def last_data_row_ws(ws, mapping, header_row=1):
    """末数据行（1-based），以**全部已映射列**有无值判定；mapping 为空返回 None。

    这是「末数据行」的**唯一判据**：`probe.analyze_workbook` 与 `next_append_row_ws`
    共用它。历史上两处各写一套（probe 用全部映射列，next_append_row 只用
    "需求描述/提出时间/提出人" 3 个探针列），于是尾部行若只在**别的**映射列
    （如 项目编号）有值，两处就会分歧：next_append_row 算出**更小**的起始行，
    `final`/`position_reuse` 便据此**静默覆盖**尾行数据（2026-09-11 实跑复现）。
    判据统一后，两处必然得出同一个起始行。

    **孤儿尾行是预期行为**：只在项目编号之类非探针列有值的尾行会被**当作有效数据行保留**
    （末行 = 它），本批从它下方追加 —— 宁可在表里留一个空行，也绝不覆盖已有单元格。
    要复用这类行需显式 `--start-row`。
    """
    if not mapping:
        return None
    cols = [col_index(l) + 1 for l in mapping.values()]
    last = header_row
    for r in range(header_row + 1, ws.max_row + 1):
        if any(ws.cell(row=r, column=c).value not in (None, "") for c in cols):
            last = r
    return last


def next_append_row_ws(ws, mapping=None, header_row=1):
    """返回工作表下一个可追加行号（1-based）。

    有 mapping 时与 probe 共用 `last_data_row_ws`（全部映射列判末行），保证
    probe 与 final/position_reuse 得出**同一个**起始行；无 mapping 时退化为
    扫描全部已有列（不再只扫前 27 列，宽表也能判准）。
    """
    last = last_data_row_ws(ws, mapping, header_row)
    if last is None:
        cols = list(range(1, max(1, ws.max_column or 0) + 1))
        last = 1
        for r in range(1, ws.max_row + 1):
            if any(ws.cell(row=r, column=c).value not in (None, "") for c in cols):
                last = r
    return last + 1


def require_openpyxl():
    """导入 openpyxl，缺库时给出可执行的提示，而不是裸 ImportError。

    只有**本地表格通道**需要它（云文档通道走快照，不需要）。全 skill 仅此一处导入，
    避免各脚本各写一份 try/except 导致提示不一致 —— 与 zstandard 那条的处理方式对齐。
    """
    try:
        import openpyxl
        return openpyxl
    except ImportError:
        raise SystemExit(
            "✗ 本地表格通道需要 openpyxl，当前解释器未安装。\n"
            "  安装：pip install openpyxl\n"
            "  （只走云文档通道的话不需要它 —— 用 --snapshot / --history 即可。）")


# ---------------- 输入路径守卫（唯一入口）----------------
# 用户把路径敲错是最高频的输入错误，而 `open()` / `load_workbook()` 抛的是
# FileNotFoundError —— 裸 traceback 对使用者毫无帮助。这几条守卫全 skill 共用一份，
# 与 require_openpyxl 同一思路：同一件事只留一个实现。
def is_placeholder(path):
    """config.example.json 里满是 `<...>` 占位符 —— 照抄没改的情形要认出来。

    判据：整串就是 `<...>`。**占位符 = 没配置**，与"配了一个错的路径"是两回事：
    前者应当按"未提供"处理（首次使用时的正常状态），后者才该报错。
    """
    s = str(path or "").strip()
    return s.startswith("<") and s.endswith(">") and len(s) > 2


def require_file(path, what="文件"):
    """路径缺失/不是文件 → SystemExit + 可执行提示（不抛裸 traceback）。"""
    if not path:
        raise SystemExit(f"✗ 未提供{what}路径。")
    if is_placeholder(path):
        raise SystemExit(f"✗ {what} 还是占位符：{path}\n"
                         "  请把 config.example.json 里的 <...> 替换成真实路径。")
    if not os.path.isfile(path):
        raise SystemExit(f"✗ {what}不存在：{path}\n"
                         "  请检查路径是否写错（相对路径按**当前工作目录**解析）。")
    return path


def require_dir(path, what="目录"):
    """目录缺失 → SystemExit + 可执行提示。"""
    if not path:
        raise SystemExit(f"✗ 未提供{what}路径。")
    if is_placeholder(path):
        raise SystemExit(f"✗ {what} 还是占位符：{path}\n"
                         "  请把 config.example.json 里的 <...> 替换成真实路径。")
    if not os.path.isdir(path):
        raise SystemExit(f"✗ {what}不存在或不是目录：{path}\n"
                         "  请检查路径是否写错（相对路径按**当前工作目录**解析）。")
    return path


def load_json_file(path, what="JSON 文件"):
    """读 JSON：路径不存在 / 不是合法 JSON 都给可执行提示。"""
    require_file(path, what)
    try:
        with open(path, encoding="utf-8") as f:
            return json.load(f)
    except json.JSONDecodeError as e:
        raise SystemExit(f"✗ {path} 不是合法 JSON（{e}）。")


def open_local_workbook(path):
    """打开本地 xlsx 工作簿，三类失败都给可执行提示。

    顺序有意为之：**先判存在与扩展名，再要 openpyxl** —— 否则用户既打错路径又没装库时，
    看到的是"请装 openpyxl"，而真正的问题是路径写错了。
    `.xls/.xlt` 是旧版二进制格式，openpyxl 不支持，与是否安装无关，所以也放在前面。
    """
    p = require_file(path, "工作簿")
    if p.lower().endswith((".xls", ".xlt")):
        raise SystemExit(f"✗ {p} 是旧版格式，openpyxl 不支持。"
                         "请先用 Excel/WPS 把工作簿「另存为 .xlsx」再运行。")
    openpyxl = require_openpyxl()
    from openpyxl.utils.exceptions import InvalidFileException
    try:
        return openpyxl.load_workbook(p)
    except InvalidFileException as e:
        raise SystemExit(f"✗ 无法以 .xlsx 解析 {p}（{e}）。"
                         "若是旧版 .xls，请先另存为 .xlsx 再运行。")


# ---------------- 表格快照（云文档 / WPS 通道）----------------
# 脚本无法直连 MCP 连接器，故两端都收敛到 JSON：
#   读：agent 调 kdocs sheet.get_range_data → 原始返回存 raw.json → sparse_to_grid 成快照
#   写：脚本产出 payload → to_kdocs_payload.py 转 rangeData → agent 调 sheet.update_range_data
# 快照结构见 sheet_snapshot.py 文档字符串。

class _Cell:
    """冒充 openpyxl 的 Cell（本 skill 只用到 .value）。"""
    __slots__ = ("value",)

    def __init__(self, value):
        self.value = value


class GridSheet:
    """用二维字符串数组冒充工作表，接口与 openpyxl 用到的部分保持一致（行列 1-based）。"""

    def __init__(self, name, rows, num_formats=None, sheet_id=None):
        self.title = name
        self.sheet_id = sheet_id
        self._rows = [[("" if v is None else str(v)) for v in r] for r in (rows or [])]
        self.num_formats = num_formats or {}

    @property
    def max_row(self):
        return len(self._rows)

    @property
    def max_column(self):
        return max((len(r) for r in self._rows), default=0)

    def cell(self, row=1, column=1):
        r, c = row - 1, column - 1
        v = ""
        if 0 <= r < len(self._rows) and 0 <= c < len(self._rows[r]):
            v = self._rows[r][c]
        return _Cell(v if v != "" else None)

    def row_values(self, row=1):
        r = row - 1
        return list(self._rows[r]) if 0 <= r < len(self._rows) else []


class GridWorkbook:
    """冒充 openpyxl Workbook（只实现 .sheetnames / __getitem__）。"""

    def __init__(self, sheets):
        self._sheets = list(sheets)
        self.sheetnames = [s.title for s in self._sheets]

    def __getitem__(self, name):
        for s in self._sheets:
            if s.title == name:
                return s
        raise KeyError(name)


def _entry_value(entry):
    """从 kdocs 单元格条目取"人看的文本"：cellText 优先（公式则退回原始值）。"""
    txt = entry.get("cellText")
    orig = entry.get("originalCellValue")
    if txt not in (None, "") and not str(txt).startswith("="):
        return str(txt)
    if orig not in (None, ""):
        return str(orig)
    return ""


def merge_range_data(store, range_data):
    """把一个 kdocs rangeData（稀疏）并入累加器 store（原地修改）。

    store = {"cells": {(r,c): str}, "num_formats": {"r,c": fmt}, "max_row": int, "max_col": int}
    """
    store.setdefault("cells", {})
    store.setdefault("num_formats", {})
    for e in range_data or []:
        if not isinstance(e, dict):
            continue
        try:
            rf = int(e.get("rowFrom", 0) or 0)
            cf = int(e.get("colFrom", 0) or 0)
            rt = int(e.get("rowTo", rf) if e.get("rowTo") is not None else rf)
            ct = int(e.get("colTo", cf) if e.get("colTo") is not None else cf)
        except (TypeError, ValueError):
            continue
        v = _entry_value(e)
        nf = e.get("numFormat")
        for r in range(min(rf, rt), max(rf, rt) + 1):
            for c in range(min(cf, ct), max(cf, ct) + 1):
                store["cells"][(r, c)] = v
                if nf:
                    store["num_formats"][f"{r},{c}"] = nf
                if r > store.get("max_row", -1):
                    store["max_row"] = r
                if c > store.get("max_col", -1):
                    store["max_col"] = c
    return store


def store_to_rows(store):
    """累加器 -> 紧凑二维列表（裁掉尾部/右侧全空行列）。"""
    cells = store.get("cells") or {}
    maxr, maxc = store.get("max_row", -1), store.get("max_col", -1)
    rows = [[cells.get((r, c), "") for c in range(maxc + 1)] for r in range(maxr + 1)]
    while rows and not any(str(x).strip() for x in rows[-1]):
        rows.pop()
    width = max((len(r) for r in rows), default=0)
    while width and all(not str(r[width - 1]).strip() for r in rows):
        width -= 1
    return [r[:width] for r in rows]


def sparse_to_grid(range_data):
    """kdocs rangeData（稀疏）-> (rows, num_formats)。"""
    store = merge_range_data({}, range_data)
    return store_to_rows(store), store.get("num_formats") or {}


def load_snapshot(path):
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def warn_if_inside_git_repo(path, what="产物"):
    """产物目录落在**某个 git 仓库内**时打醒目告警；返回是否命中。

    起因：`rules_check.py plan` 的默认 `--out` 是 `<cwd>/rules_check_out`，从仓库根目录跑一次，
    就把**含真实标识**的产物（渲染后的提示词里有我方标识=微信号等）写进了仓库 —— 而未跟踪、
    也未忽略的目录**不会出现在 `.gitignore` 的检查里**，`git add -A` 会顺手带走。
    `.gitignore` 只能堵住**已知名字**，这里补一道"落到仓库里就说一声"，给新产物兜底。

    刻意做成**非致命**：git 不存在 / 不是仓库 / 跨盘符 → 静默返回 False。
    告警机制本身失败，绝不能影响主流程。
    """
    try:
        want = os.path.abspath(path)            # 用户**想写**的落点（可能还不存在）
        d = want if os.path.isdir(want) else (os.path.dirname(want) or os.getcwd())
        pr = subprocess.run(["git", "-C", d, "rev-parse", "--show-toplevel"],
                            capture_output=True, timeout=15)
        if pr.returncode != 0:
            return False
        root = os.path.abspath(pr.stdout.decode("utf-8", "replace").strip())
        if not root:
            return False
        # ⚠ Windows 上 `TEMP` 常是 **8.3 短名**（形如 `...~1\AppData\Local\Temp`），
        #   而 `git rev-parse --show-toplevel` 把仓库根归一成**长名** ——
        #   两者字符串毫无共同前缀可言，于是 `commonpath` 会算成一个很浅的祖先目录
        #   而不是仓库根，"在仓库内"被误判为"不在"，**告警静默失效**。
        #   真事故：CI 的 `windows-latest` 上这条一直不生效，而本地 `TEMP` 恰好是长名
        #   ⇒ 永远复现不了。`realpath` 会把短名展开成长名（Windows 上即 `GetLongPathName`），
        #   两边都归一后再比；再用 `normcase` 抹掉盘符大小写差异。
        root = os.path.realpath(root)
        d = os.path.realpath(d)
        try:
            if os.path.normcase(os.path.commonpath([root, d])) != os.path.normcase(root):
                return False
        except ValueError:                      # 不同盘符 → 必然不在同一仓库
            return False
    except Exception:                           # 没装 git / 超时 / 权限
        return False
    rel = os.path.relpath(want, root).replace(os.sep, "/")
    print(f"⚠ {what}落在 git 仓库内（{root} 下的 {rel}）：内容含真实标识，**不要入库**。\n"
          "  请确认 .gitignore 已覆盖它；或把 --out 指到仓库外的目录。")
    return True


def snapshot_sheet_coverage(snap, sheet_name=None):
    """取快照里某子表「本次读表**实际覆盖**到的行列范围」（0-based，含端点）。

    为什么需要它：kdocs 回包**不带请求范围**（只有稀疏的 rangeData，没值的格根本不出现），
    所以快照无法自证"我请求了多大范围"，只能自证"我看到了哪里"。
    `sheet_snapshot.py build` 会把累加器的原始极值（**裁边之前**）写进 `coverage` ——
    这正是判断"本次读表有没有覆盖到数据区 / 岗位列"的判据（裁边后的 max_row/max_column
    会被尾部空行、右侧空列缩小，用它判断会误报）。

    找不到 coverage（旧快照 / 手工构造的）返回 None，调用方据此跳过检查、不误报。
    """
    for s in (snap or {}).get("sheets") or []:
        if sheet_name is not None and (s.get("name") or "") != sheet_name:
            continue
        cov = s.get("coverage")
        return cov if isinstance(cov, dict) else None
    return None


def workbook_from_snapshot(snap):
    """快照 dict -> GridWorkbook。"""
    sheets = []
    for i, s in enumerate((snap or {}).get("sheets") or [], 1):
        rows = s.get("rows")
        if rows is None:
            rows, fmts = sparse_to_grid(s.get("rangeData"))
            s.setdefault("num_formats", fmts)
        sheets.append(GridSheet(s.get("name") or f"sheet{i}", rows,
                                s.get("num_formats"),
                                s.get("worksheet_id", s.get("sheetId"))))
    if not sheets:
        raise ValueError("快照里没有任何工作表（sheets 为空）：请先按 SKILL.md 的"
                         "「WPS 通道读表」把 kdocs 返回存成 raw.json 再 build。")
    return GridWorkbook(sheets)


# ---------------- 岗位列复用（final 与 position_reuse 共用这一套）----------------
# 设计要点：读历史、填新行、行序传播三件事各只有一份实现，
# 避免"final 内置"与"position_reuse 独立"两条路各写一套而逐渐走偏。

def position_columns(mapping):
    """从 column_mapping 解析 (提出人列, 岗位列)（1-based）。

    缺映射时退回常见列位（提出人=O=15, 岗位=N=14）。返回 (col_person, col_post)。
    """
    mp = mapping or {}
    col_person = col_index(mp["提出人"]) + 1 if mp.get("提出人") else 15
    col_post = col_index(mp["提出人岗位"]) + 1 if mp.get("提出人岗位") else 14
    return col_person, col_post


def read_history_positions(sheet, header_row, end_row, col_person, col_post):
    """读历史区的「提出人 -> 该人最近的非空岗位」。

    **只读**：历史区永不修改，只作为新行的岗位来源。
    区间为 [header_row+1, end_row]（1-based，含两端）；end_row 一般传「追加起始行-1」。

    返回 (last_post, pairs)；pairs 是读到「人+岗位」成对的行数——为 0 时说明
    表格来源没覆盖到这两列的数据区（快照常见），调用方应据此告警而不是静默补 0 条。
    """
    last_post = {}
    pairs = 0
    for r in range(int(header_row) + 1, int(end_row) + 1):
        person = sheet.cell(row=r, column=col_person).value
        if person is None or not str(person).strip():
            continue
        person = str(person).strip()
        post = sheet.cell(row=r, column=col_post).value
        if post is not None and str(post).strip():
            last_post[person] = str(post).strip()
            pairs += 1
    return last_post, pairs


def reuse_position_fill(new_rows, last_post, person_key="提出人", post_key="提出人岗位",
                        override=False):
    """把 last_post（人 -> 最近岗位）复用到 new_rows 的空岗位上，**原地修改** new_rows。

    返回变更明细 [{"i": 行下标, "person": 人, "old": 原值, "new": 补的值}]。

    语义（与早期 position_reuse 完全一致，只是挪到共用函数里）：
      - 默认**只填空缺**；override=True 才允许把已有值改写成 last_post 里的值。
      - 按行序推进：本批先出现的非空岗位会写回 last_post，供本批后续行使用，
        因此"历史 + 本批"合起来就是整表从上到下的自然语义。
      - last_post 会被原地更新（调用方若还要用需自己拷贝）。
    """
    changes = []
    for i, r in enumerate(new_rows or []):
        person = str(r.get(person_key) or "").strip()
        if not person:
            continue
        cur = str(r.get(post_key) or "").strip()
        hist = str(last_post.get(person) or "").strip()
        if hist and ((not cur) or (override and cur != hist)):
            changes.append({"i": i, "person": person, "old": cur, "new": hist})
            r[post_key] = hist
            cur = hist
        if cur:
            last_post[person] = cur
    return changes


if __name__ == "__main__":
    ensure_utf8_stdio()
    cfg, path = load_config()
    print("config:", path or "(未找到 config.json，请先从 config.example.json 复制)")
    print("db_storage:", find_db_storage())
    print("wcdb_tool :", find_wcdb_tool() or "(未找到，请 clone 或设 WCDB_KEY_TOOL)")
