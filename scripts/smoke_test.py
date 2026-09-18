# -*- coding: utf-8 -*-
"""wechat-worklog-matrix 冒烟自检：import 全模块 + 纯函数断言 + CLI 契约 + 离线端到端。

用法：
  python scripts/smoke_test.py            # 无第三方依赖的部分（CI 的 deps=minimal 格子跑这个）
  python scripts/smoke_test.py --full     # 额外跑依赖 openpyxl 的部分

不触碰真实微信数据/网络，可随时执行。
`.github/workflows/selftest.yml` 跑的就是它 —— 本地与 CI **同一个入口**，
避免出现"CI 绿、本地跑不到"或反之。
"""
import sys
import os
import re
import fnmatch
import subprocess

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, HERE)

FAIL = []
STAT = {"ok": 0, "skip": 0}


class Skip(Exception):
    """本环境不适用的检查：显式标 `[SKIP]`。

    刻意**不**把跳过算成通过 —— 本项目的基本原则是"绝不静默"，自检自身也不能例外
    （否则会出现"绿着但根本没跑"）。"""
    pass


def check(name, fn, count=True):
    """跑一项检查。`count=False` 用于"核对项数"这类**不该把自己算进去**的元检查。"""
    try:
        fn()
        if count:
            STAT["ok"] += 1
        print(f"  [OK] {name}")
    except Skip as e:
        if count:
            STAT["skip"] += 1
        print(f"  [SKIP] {name} —— {e}")
    except Exception as e:
        FAIL.append(name)
        print(f"  [FAIL] {name}: {type(e).__name__}: {e}")


MODULES = ["common", "probe", "decrypt", "export_conversations", "split_for_agents",
           "build_matrix_rows", "position_reuse", "sheet_snapshot", "to_kdocs_payload"]


def smoke_imports():
    import importlib
    for m in MODULES:
        importlib.import_module(m)


def smoke_common():
    from common import serial, col_index, col_letter
    assert serial("2026-05-14") == 46156
    assert serial("2026/5/14") == 46156
    assert serial("2026年5月14日") == 46156
    assert serial("2026-05-14 10:00") == 46156
    assert serial("bad") == ""
    assert col_index("A") == 0 and col_index("AA") == 26
    assert col_letter(0) == "A" and col_letter(26) == "AA"


def smoke_snapshot_grid():
    """快照抽象层：稀疏 kdocs 返回 -> 密集网格，且接口与 openpyxl 一致。"""
    from common import (sparse_to_grid, merge_range_data, store_to_rows,
                        GridSheet, GridWorkbook, workbook_from_snapshot,
                        next_append_row_ws, pick_sheet)
    rd = [
        {"rowFrom": 0, "rowTo": 0, "colFrom": 0, "colTo": 0, "cellText": "需求描述"},
        {"rowFrom": 0, "rowTo": 0, "colFrom": 1, "colTo": 1, "cellText": "提出时间", "numFormat": "yyyy-mm-dd"},
        {"rowFrom": 0, "rowTo": 0, "colFrom": 2, "colTo": 2, "cellText": "提出人"},
        {"rowFrom": 1, "rowTo": 1, "colFrom": 0, "colTo": 0, "cellText": "改个费用"},
        {"rowFrom": 1, "rowTo": 1, "colFrom": 1, "colTo": 1, "originalCellValue": "46156",
         "cellText": "2026-05-14", "numFormat": "yyyy-mm-dd"},
        {"rowFrom": 1, "rowTo": 1, "colFrom": 2, "colTo": 2, "cellText": "张三"},
    ]
    rows, fmts = sparse_to_grid(rd)
    assert rows[0] == ["需求描述", "提出时间", "提出人"], rows
    assert rows[1][1] == "2026-05-14", rows[1]        # cellText 优先于原始序列号
    assert fmts.get("0,1") == "yyyy-mm-dd"
    # 公式单元格退回原始值
    rows2, _ = sparse_to_grid([{"rowFrom": 0, "rowTo": 0, "colFrom": 0, "colTo": 0,
                                "cellText": "=A1", "originalCellValue": "=A1"}])
    assert rows2[0][0] == "=A1", rows2
    # 稀疏：中间空单元格应被补空
    rows3, _ = sparse_to_grid([{"rowFrom": 0, "rowTo": 0, "colFrom": 0, "colTo": 0, "cellText": "a"},
                               {"rowFrom": 0, "rowTo": 0, "colFrom": 2, "colTo": 2, "cellText": "c"}])
    assert rows3[0] == ["a", "", "c"], rows3
    # range 覆盖（合并单元格语义）：整块填同一个值
    store = merge_range_data({}, [{"rowFrom": 0, "rowTo": 0, "colFrom": 0, "colTo": 3,
                                   "cellText": "x"}])
    assert store_to_rows(store) == [["x", "x", "x", "x"]], store_to_rows(store)
    # 尾部空行应被裁掉（range 覆盖到空单元格时）
    store2 = merge_range_data({}, [{"rowFrom": 0, "rowTo": 0, "colFrom": 0, "colTo": 0,
                                    "cellText": "x"},
                                   {"rowFrom": 1, "rowTo": 5, "colFrom": 0, "colTo": 0,
                                    "cellText": ""}])
    assert store_to_rows(store2) == [["x"]], store_to_rows(store2)
    # GridShell 接口与 openpyxl 同形
    gs = GridSheet("运维-2026", rows)
    assert gs.max_row == 2 and gs.max_column == 3
    assert gs.cell(row=1, column=1).value == "需求描述"
    assert gs.cell(row=9, column=9).value is None      # 越界 -> None（同 openpyxl 空单元格）
    wb = GridWorkbook([gs, GridSheet("说明", [])])
    assert wb.sheetnames == ["运维-2026", "说明"]
    assert pick_sheet(wb, "运维").title == "运维-2026"
    assert pick_sheet(wb, "不存在").title == "运维-2026"
    # 追加起始行：有映射时按"需求描述/提出时间/提出人"列判定
    assert next_append_row_ws(gs, {"需求描述": "A"}) == 3
    assert next_append_row_ws(GridSheet("x", []), {"需求描述": "A"}) == 2
    # 由快照 dict 还原
    wb2 = workbook_from_snapshot({"sheets": [{"name": "运维", "sheetId": 7, "rows": rows}]})
    assert wb2["运维"].sheet_id == 7
    try:
        workbook_from_snapshot({"sheets": []})
    except ValueError:
        pass
    else:
        raise AssertionError("空快照应抛 ValueError")


def smoke_snapshot_raw_parse():
    """raw.json 的多形态识别（kdocs 返回包装不一，必须都吃得下）。"""
    from sheet_snapshot import extract_range_data
    cell = [{"rowFrom": 0, "colFrom": 0, "cellText": "a"}]
    assert extract_range_data({"result": "ok", "detail": {"rangeData": cell}}) == cell
    assert extract_range_data({"rangeData": cell}) == cell
    assert extract_range_data(cell) == cell
    assert extract_range_data({"data": {"rangeData": cell}}) == cell
    assert extract_range_data({"detail": {"detail": {"rangeData": cell}}}) == cell
    import json as _json
    wrapped = {"content": [{"type": "text", "text": _json.dumps({"detail": {"rangeData": cell}})}]}
    assert extract_range_data(wrapped) == cell
    assert extract_range_data({"code": 0}) == []
    assert extract_range_data(None) == []


def smoke_kdocs_payload():
    """payload -> kdocs rangeData 转换（editor_sdk 格式 / 简化格式 / 日期列 numfmt）。"""
    from to_kdocs_payload import (normalize_values, build_range_data, build_body,
                                  _num_to_str, _runs)
    a = {"values": [{"row": 1, "col": 11, "value_type": "STRING", "string_value": "改费"},
                    {"row": 1, "col": 12, "value_type": "NUMBER", "number_value": 46156},
                    {"row": 1, "col": 13, "value_type": "STRING", "string_value": ""}]}
    cells = normalize_values(a)
    assert cells == [(1, 11, "改费"), (1, 12, "46156")], cells   # 空值被剔除
    b = {"values": [{"row": 0, "col": 0, "value": "x"}]}
    assert normalize_values(b) == [(0, 0, "x")]
    assert _num_to_str(46156.0) == "46156" and _num_to_str(1.5) == "1.5"
    assert _runs([1, 2, 3, 7, 8]) == [(1, 3), (7, 8)] and _runs([]) == []
    rd = build_range_data(cells, date_cols0=[12])
    assert rd[0] == {"opType": "formula", "rowFrom": 1, "rowTo": 1, "colFrom": 11,
                     "colTo": 11, "formula": "改费"}, rd[0]
    assert rd[-1]["opType"] == "format" and rd[-1]["xf"]["numfmt"] == "yyyy/m/d", rd
    # 同列连续行压成一段
    nhi = build_range_data([(0, 12, "1"), (1, 12, "2"), (2, 12, "3")], date_cols0=[12])
    fmts = [x for x in nhi if x["opType"] == "format"]
    assert len(fmts) == 1 and fmts[0]["rowFrom"] == 0 and fmts[0]["rowTo"] == 2, nhi
    # 不补格式时只返回值
    assert all(x["opType"] == "formula" for x in build_range_data(cells, [12], with_date_format=False))

    # D3: 「会被引擎改写」的纯文本必须能被检出并**无损转义**（真机实测：前置单引号
    # 被引擎当文本标记消费掉、回读值不含引号），且不得误报普通业务文本 / 日期序列号。
    from to_kdocs_payload import find_risky_text, risky_reason, escape_risky_text
    risky = find_risky_text([(0, 0, "0012"), (0, 1, "=A1"), (0, 2, "12345678901234567"),
                             (0, 3, "+86"), (0, 4, "正常需求描述"), (0, 5, "2026-09-11"),
                             (0, 6, "-"), (0, 7, "商务经理")])
    assert [r[2] for r in risky] == ["0012", "=A1", "12345678901234567", "+86"], risky
    assert find_risky_text([]) == []
    assert risky_reason("46274") is None          # 日期序列号必须保持"可被转成数字"
    assert risky_reason("=") is None and risky_reason("-") is None   # 单字符真机实测不改写
    assert escape_risky_text("0012") == "'0012"
    assert escape_risky_text("商务经理") == "商务经理"
    assert escape_risky_text("46274") == "46274"
    assert escape_risky_text("'abc") == "''abc"    # 本身带引号 → 再补一层，存回来仍是 'abc
    # 关键契约：工作表键名必须是 worksheet_id（连接器参数名），不是 sheetId
    body = build_body("F1", 3, rd)
    assert set(body) == {"file_id", "worksheet_id", "rangeData"}, body.keys()
    assert body["worksheet_id"] == 3 and "sheetId" not in body
    assert build_body(None, None, [])["worksheet_id"] == "<worksheet_id>"
    # 关键契约：单次 rangeData 上限 100 条 → 必须自动分批
    from to_kdocs_payload import chunk_range_data
    big = [{"opType": "formula"}] * 110
    ch = chunk_range_data(big, 100)
    assert [len(c) for c in ch] == [100, 10], [len(c) for c in ch]
    assert sum(len(c) for c in ch) == 110
    assert chunk_range_data([], 100) == [[]]
    assert [len(c) for c in chunk_range_data(big, 1000)] == [110]


def smoke_probe_analyze():
    """probe.analyze_workbook 应对 GridWorkbook（云文档快照）同样可用。"""
    from common import GridSheet, GridWorkbook
    from probe import analyze_workbook
    rows = [
        ["项目编号", "项目名称", "需求描述", "提出时间", "提出人岗位", "提出人", "解决人"],
        ["P1", "某项目", "改个费用", "2026-05-14", "客服", "张三", "李四"],
        ["", "", "导数据", "2026-05-15", "", "张三", "李四"],
    ]
    res = analyze_workbook(GridWorkbook([GridSheet("运维-2026", rows)]), None, "snap.json")
    assert res["header_row"] == 1
    assert res["column_mapping"]["需求描述"] == "C", res["column_mapping"]
    assert res["column_mapping"]["提出人"] == "F"
    assert res["column_mapping"]["解决人"] == "G"
    assert res["next_append_row"] == 4, res["next_append_row"]
    assert res["date_columns"] == ["D"], res["date_columns"]
    assert res["handler_candidates"][0]["value"] == "李四", res["handler_candidates"]

    # N3: 表头行的主判据是「命中内置别名数」，不能被"更满的数据行"骗过
    #（种子表：数据行 4 个非空 > 表头 3 个非空，但别名命中 0 < 3）。
    # 旧规则只比非空数 → 会把数据行当表头 → 列映射为空（虽然会被 C1 守卫拦住，但整表读不了）。
    sloppy = [
        ["项目编号", "需求描述", "提出人"],
        ["P1", "改个费用", "张三", "多出来的一列"],
    ]
    res2 = analyze_workbook(GridWorkbook([GridSheet("运维", sloppy)]), None, "snap.json")
    assert res2["header_row"] == 1, res2["header_row"]
    assert res2["column_mapping"].get("需求描述") == "B", res2["column_mapping"]
    assert res2["column_mapping"].get("提出人") == "C", res2["column_mapping"]


def smoke_message_decoding():
    """消息正文解码：ZSTD 压缩 / appmsg 引用 / 系统消息 归一化。"""
    from export_conversations import (render_content, decode_payload, appmsg_summary,
                                      UNDECODED_TAG, BINARY_TAG, ZSTD_MAGIC)
    # zstd 帧但解不开（测试环境无 zstandard 或帧非法）→ 明确标记，绝不吐乱码
    bad = ZSTD_MAGIC + b"\x00" * 16
    assert render_content(bad, 4, 1) == UNDECODED_TAG, render_content(bad, 4, 1)
    # 纯二进制且高乱码率 → 标记
    assert decode_payload(b"\xff\xfe\x00\x01" * 20, 0)[1] == BINARY_TAG
    # 明文直通
    assert render_content("你好", 0, 1) == "你好"
    # 媒体占位（优先级高于内容）
    assert render_content("x", 0, 3) == "[图片]"
    assert render_content("x", 0, 34) == "[语音]"
    # appmsg：提取标题 + 引用正文（对方用「引用回复」提需求时正文在引用里）
    xml = ('<?xml version="1.0"?><msg><appmsg><title>话头示例A</title>'
           '<refermsg><content>引用正文示例B</content></refermsg>'
           '</appmsg></msg>')
    got = render_content(xml, 0, 49)
    assert got.startswith("[链接/文件] ") and "话头示例A" in got and "引用正文示例B" in got, got
    assert appmsg_summary("<msg></msg>") == ""
    # N4 兜底：无 title/des/引用时，外层 <content> 或 <url> 也要能取出来，否则只剩 `[链接/文件]`
    assert appmsg_summary('<msg><appmsg><content>纯文本示例C</content></appmsg></msg>') \
        == "纯文本示例C"
    assert appmsg_summary('<msg><appmsg><url>https://x.cn/a</url></appmsg></msg>') == "https://x.cn/a"
    # 引用正文只能算一次：有 refermsg 时兜底不得把它再抓一遍
    only_ref = '<msg><appmsg><refermsg><content>引用正文</content></refermsg></appmsg></msg>'
    assert appmsg_summary(only_ref) == "引用: 引用正文", appmsg_summary(only_ref)
    # 系统消息 / 撤回 / 通话
    assert render_content('<sysmsg type="revokemsg"><content>x</content></sysmsg>', 0, 10000) == "[撤回了一条消息]"
    assert render_content('<sysmsg type="roomtoolstips">x</sysmsg>', 0, 10000) == "[系统消息]"
    assert render_content('<voipmsg><msg>通话</msg></voipmsg>', 0, 50) == "[通话]"
    # 高位 local_type 要按 32 位掩码还原
    assert render_content("x", 0, 244813135921) == "x", "掩码后应等于 49(appmsg) 且正文直通"
    assert render_content(xml, 0, (49 + (57 << 32))) == got


def smoke_guards_and_robustness():
    """真跑暴露的守卫：空映射不给追加行、岗位剥离不剥空、空 N 键不丢、null 日期不崩。"""
    import json as _json
    import tempfile
    from common import GridSheet, GridWorkbook
    from probe import analyze_workbook
    from build_matrix_rows import norm_o, load_agents, flag_dups

    # C1: 表头全部识别失败 → 不能给 next_append_row（否则会覆盖第 2 行）
    wb = GridWorkbook([GridSheet("运维", [["甲", "乙"], ["1", "2"], ["3", "4"]])])
    res = analyze_workbook(wb, None, "snap.json")
    assert res["column_mapping"] == {}
    assert res["last_data_row"] is None and res["next_append_row"] is None, res
    assert res["warnings"], "空映射必须给出 warning"
    # 有映射时仍正常
    wb2 = GridWorkbook([GridSheet("运维", [["需求描述", "提出人"], ["a", "张三"], ["b", "李四"]])])
    r2 = analyze_workbook(wb2, None, "snap.json")
    assert r2["next_append_row"] == 4, r2["next_append_row"]

    # C9: 无映射时扫描全部列（不再只扫前 27 列）
    wide = GridSheet("w", [[""] * 40 + ["x"]])
    from common import next_append_row_ws
    assert next_append_row_ws(wide, None) == 2, next_append_row_ws(wide, None)

    # N5: probe 与 final/position_reuse 的「追加起始行」必须**同源同值**。
    # 尾行只在"非 3 探针列"的映射列有值时，旧实现会算出更小的起始行 → 覆盖尾行。
    mp = {"项目编号": "A", "需求描述": "B", "提出时间": "C", "提出人": "D"}
    tail_only_pid = GridWorkbook([GridSheet("运维", [
        ["项目编号", "需求描述", "提出时间", "提出人"],
        ["P1", "历史甲", "2026-09-01", "张三"],
        ["P2", "历史乙", "2026-09-02", "李四"],
        ["P3", "", "", ""],                       # 尾行只在项目编号列有值
    ])])
    _a = analyze_workbook(tail_only_pid, None, "snap.json")
    _b = next_append_row_ws(tail_only_pid["运维"], _a["column_mapping"])
    assert _a["next_append_row"] == _b == 5, (_a["next_append_row"], _b)
    # 表头不在第 1 行时，两处也要一致（header_row 必须被透传）
    two_hdr = GridWorkbook([GridSheet("运维", [
        ["2026年运维台账", "", ""],
        ["项目编号", "需求描述", "提出人"],
        ["P1", "历史甲", "张三"],
    ])])
    _c = analyze_workbook(two_hdr, None, "snap.json")
    assert _c["header_row"] == 2, _c["header_row"]
    assert _c["next_append_row"] == 4, _c["next_append_row"]
    assert next_append_row_ws(two_hdr["运维"], _c["column_mapping"], 2) == 4

    # C10: 姓名恰好等于岗位词时不剥空
    cfg = {"post_words": ["客服", "财务"]}
    assert norm_o("客服", cfg) == ("客服", ""), norm_o("客服", cfg)
    assert norm_o("张三客服", cfg) == ("张三", "客服")

    # C6/C7: ask_date=null 不崩；N 键存在但为空时不丢剥离结果
    d = tempfile.mkdtemp(prefix="wm_smoke_")
    _json.dump([{"O": "王五客服", "N": "", "L": "改费", "ask_date": "2026-05-14", "chat": "A"},
                {"O": "李四", "L": "导数据", "ask_date": None, "chat": "B"}],
               open(os.path.join(d, "agent_1.json"), "w", encoding="utf-8"), ensure_ascii=False)
    rows = load_agents(d, cfg)
    assert len(rows) == 2
    by_o = {r["O"]: r for r in rows}
    assert by_o["王五"]["N"] == "客服", by_o["王五"]
    assert rows[0]["ask_date"] is None or isinstance(rows[0]["ask_date"], str)

    # C8: 日期不可解析时不崩、不误判
    flag_dups([{"O": "a", "L": "改个费用", "ask_date": "坏日期", "chat": "A", "_i": 1},
               {"O": "a", "L": "改个费用", "ask_date": "2026-05-14", "chat": "B", "_i": 2}], {})
    import shutil
    shutil.rmtree(d, ignore_errors=True)


def smoke_reuse_position():
    """岗位复用共用层：读历史只到 end_row、默认只填空缺、批内顺序传播。"""
    from common import GridSheet, position_columns, read_history_positions, reuse_position_fill

    # 列解析：有映射按映射，缺映射退回 提出人=O(15) / 岗位=N(14)
    assert position_columns({"提出人": "O", "提出人岗位": "N"}) == (15, 14)
    assert position_columns({"提出人": "C", "提出人岗位": "B"}) == (3, 2)
    assert position_columns({}) == (15, 14)

    gs = GridSheet("运维", [
        ["需求描述", "提出人岗位", "提出人"],   # 第 1 行 = 表头
        ["a", "客服", "张三"],                  # 第 2 行
        ["b", "", "张三"],                      # 第 3 行
        ["c", "", "李四"],                      # 第 4 行（李四始终没岗位）
        ["d", "商务", "张三"],                  # 第 5 行
    ])
    # 只读到 end_row（历史区右端 = 追加起始行-1）→ 读到的是"客服"
    hist, pairs = read_history_positions(gs, 1, 3, col_person=3, col_post=2)
    assert hist == {"张三": "客服"} and pairs == 1, (hist, pairs)
    # 扩到第 5 行：同一人取**最近**的非空值
    hist2, pairs2 = read_history_positions(gs, 1, 5, col_person=3, col_post=2)
    assert hist2 == {"张三": "商务"} and pairs2 == 2, (hist2, pairs2)
    # 只读：网格没被改
    assert gs.cell(row=3, column=3).value == "张三"

    # 默认只填空缺；历史里没有的人不补；无提出人的行跳过
    rows = [{"提出人": "张三", "提出人岗位": ""},
            {"提出人": "张三", "提出人岗位": "财务"},
            {"提出人": "王五", "提出人岗位": ""},
            {"提出人": "", "提出人岗位": ""}]
    ch = reuse_position_fill(rows, {"张三": "客服"}, override=False)
    assert len(ch) == 1 and ch[0]["i"] == 0 and ch[0]["new"] == "客服", ch
    assert rows[0]["提出人岗位"] == "客服"
    assert rows[1]["提出人岗位"] == "财务", rows[1]     # 默认不覆盖已有值
    assert rows[2]["提出人岗位"] == ""                  # 历史无此人
    assert rows[3]["提出人岗位"] == ""

    # override=True 才覆盖
    rows2 = [{"提出人": "张三", "提出人岗位": "财务"},
             {"提出人": "张三", "提出人岗位": ""}]
    ch2 = reuse_position_fill(rows2, {"张三": "客服"}, override=True)
    assert [c["i"] for c in ch2] == [0, 1], ch2
    assert rows2[0]["提出人岗位"] == "客服"

    # 批内顺序传播：本批前一行已有的岗位供本批后续行使用（历史为空也能补）
    rows3 = [{"提出人": "赵六", "提出人岗位": "调度"},
             {"提出人": "赵六", "提出人岗位": ""}]
    ch3 = reuse_position_fill(rows3, {}, override=False)
    assert len(ch3) == 1 and ch3[0]["i"] == 1 and ch3[0]["new"] == "调度", ch3


def smoke_build():
    from build_matrix_rows import norm_o, ids, QMAP
    cfg = {"people": {"counterparty_keyword": "李四"}, "post_words": ["财务", "客服"]}
    assert norm_o("李四财务", cfg) == ("李四", "财务"), norm_o("李四财务", cfg)
    cfg2 = {"people": {"counterparty_keyword": "某公司"}, "post_words": ["客服"]}
    assert norm_o("某公司 王五客服", cfg2) == ("王五", "客服")
    assert "ANTSQCY250789019" in ids("单号 ANTSQCY250789019 改费")
    assert "20260514" not in ids("20260514 提了需求")
    assert QMAP  # 非空


def smoke_position_reuse():
    import position_reuse      # noqa: F401


def smoke_position_reuse_openpyxl():
    """--full：岗位复用需要 openpyxl 的部分（用临时网格走快照通道，不碰真实文件）。"""
    import position_reuse      # noqa: F401
    from common import GridSheet, GridWorkbook
    gs = GridSheet("运维", [
        ["需求描述", "提出人岗位", "提出人"],
        ["a", "客服", "张三"],
        ["b", "", "张三"],
    ])
    assert gs.max_row == 3 and gs.cell(row=2, column=2).value == "客服"


def smoke_rule_consistency():
    """判定规则的三处一致性（防漂移）。

    背景：`outcome` 在预览 CSV 里靠 `OUTCOME_LABEL` 翻中文，标签缺失只会**静默**显示英文原值；
    规则开关同时出现在「config.example.json」「提示词开关渲染表」「SKILL.md」三处，
    改名时漏改一处就会让执行代理渲染出错的提示词。两者都不抛异常，只能靠断言守。
    """
    import json
    import re
    from build_matrix_rows import OUTCOME_LABEL

    root = os.path.dirname(HERE)
    prompt = open(os.path.join(root, "references", "agent-prompt-zh.txt"), encoding="utf-8").read()
    skill = open(os.path.join(root, "SKILL.md"), encoding="utf-8").read()
    cfg = json.load(open(os.path.join(root, "config.example.json"), encoding="utf-8"))

    m = re.search(r'"outcome"\s*:\s*"([^"]*)"', prompt)
    assert m, "提示词模板里找不到 outcome 枚举"
    enums = m.group(1).split("|")
    assert len(enums) == len(set(enums)), f"outcome 枚举有重复: {enums}"
    missing = [e for e in enums if e not in OUTCOME_LABEL]
    assert not missing, f"OUTCOME_LABEL 缺少 {missing}（预览会显示英文原值）"

    rules = cfg.get("rules") or {}
    prompt_switches = set(re.findall(r"rules\.([a-z_]+)", prompt))
    for k in ("ignore_if_rejected", "ignore_if_no_reply",
              "ignore_vague_complaints", "data_change_default_done"):
        assert k in rules, f"config.example.json 缺 rules.{k}"
        assert k in prompt_switches, f"提示词开关渲染表缺 rules.{k}"
        assert k in skill, f"SKILL.md 未提到 rules.{k}"

    # 「笼统抱怨」规则的核心措辞必须真的在提示词里（不能只写在文档里）
    assert "笼统抱怨" in prompt and "运维性动作与安抚不算完成" in prompt


def smoke_openpyxl_hint():
    """缺 openpyxl 时必须给可执行的提示，而不是裸 ImportError。

    用 meta_path 阻断器模拟"未安装"（不依赖真实卸载），测完恢复现场。
    背景：本地表格通道的懒加载曾直接 `import openpyxl`，缺库抛裸 traceback，
    而 zstandard 那条已有友好提示 —— 两处不对称，已收敛到 common.require_openpyxl。
    """
    import importlib.abc
    from common import require_openpyxl

    class _Block(importlib.abc.MetaPathFinder):
        def find_spec(self, name, path=None, target=None):
            if name == "openpyxl" or name.startswith("openpyxl."):
                raise ImportError("blocked for test")
            return None

    blocker = _Block()
    saved = {k: v for k, v in sys.modules.items()
             if k == "openpyxl" or k.startswith("openpyxl.")}
    for k in saved:
        del sys.modules[k]
    sys.meta_path.insert(0, blocker)
    try:
        try:
            require_openpyxl()
            raise AssertionError("缺 openpyxl 时未报错")
        except SystemExit as e:
            msg = str(e)
            assert "pip install openpyxl" in msg, msg
            assert "云文档通道" in msg, msg
    finally:
        sys.meta_path.remove(blocker)
        sys.modules.update(saved)

    # 真装了的话，正常路径应返回模块本身
    try:
        import openpyxl
        assert require_openpyxl().__name__ == "openpyxl"
    except ImportError:
        pass


# ---------------- CLI 契约 / 离线端到端 / 仓库卫生 / 控制台编码 ----------------
def _run(argv, cwd=None, env=None, timeout=180):
    """跑一个子进程，返回 (returncode, stdout, stderr)，统一按 UTF-8 解码。

    子进程若没在输出前调用 `ensure_utf8_stdio()`，在 Windows 本地 locale（cp936）下
    写中文到管道就会 `UnicodeEncodeError` → 返回码非 0 → 断言失败。这正是我们要守的。
    """
    pr = subprocess.run(argv, cwd=cwd, env=env, capture_output=True, timeout=timeout)
    return (pr.returncode,
            pr.stdout.decode("utf-8", "replace"),
            pr.stderr.decode("utf-8", "replace"))


def _mk_offline_fixture(d):
    """造一套**完全离线**的夹具：快照 + config + 裁决后的 preview CSV。

    刻意走**快照通道**（`--snapshot`）—— 它不需要 openpyxl，于是 CI 的
    「零依赖」与「装了依赖」两种格子里跑的是**同一条**断言，覆盖面对齐。
    """
    import json as _json
    import csv as _csv
    from build_matrix_rows import PREVIEW_COLS

    # final 依赖这几个键，缺一个就会 KeyError —— 先钉住这个契约
    assert {"序号", "提出人", "提出时间", "需求描述"} <= set(PREVIEW_COLS), PREVIEW_COLS

    # 一张 6 列小台账：表头第 1 行，历史区只有第 2 行（张三 / 客服）
    rows = [["项目编号", "需求描述", "提出时间", "提出人岗位", "提出人", "解决人"],
            ["P1", "历史甲", "2026-09-01", "客服", "张三", "李四"]]
    snap = os.path.join(d, "snap.json")
    with open(snap, "w", encoding="utf-8") as f:
        _json.dump({"sheets": [{"name": "运维", "sheetId": 7, "rows": rows}]}, f,
                   ensure_ascii=False)

    cfg = _json.load(open(os.path.join(ROOT, "config.example.json"), encoding="utf-8"))
    cfg["excel"].update({
        "source": "local",
        "header_row": 1,
        # 清掉 example 里的 <...> 占位符：否则"不给表格来源"的用例会先撞上占位符检查，
        # 测不到它本来要守的那条守卫（"拿不到起始行就报错退出"）
        "workbook": None,
        "start_row": None,
        # 只覆盖列映射：本批用的是"另一套模板"（A~F），列位置依然靠表头名识别
        "column_mapping": {"项目编号": "A", "需求描述": "B", "提出时间": "C",
                           "提出人岗位": "D", "提出人": "E", "解决人": "F"},
        "date_columns": ["C"],
    })
    cfg_path = os.path.join(d, "config.json")
    with open(cfg_path, "w", encoding="utf-8") as f:
        _json.dump(cfg, f, ensure_ascii=False)

    prev = os.path.join(d, "preview.csv")
    with open(prev, "w", newline="", encoding="utf-8-sig") as f:
        w = _csv.writer(f)
        w.writerow(PREVIEW_COLS)
        w.writerow(["1", "会话A", "张三", "", "2026-09-05", "把结算单的费用改成 1200",
                    "优化或需求", "中", "", "答复完成", "2026-09-06", "2026-09-06", "", "", ""])
        w.writerow(["2", "会话A", "张三", "", "2026-09-06", "搜索箱号要等 15 秒",
                    "bug", "高", "", "答复完成", "2026-09-07", "2026-09-07", "", "", ""])
    return {"snap": snap, "config": cfg_path, "preview": prev, "out": os.path.join(d, "_out")}


def smoke_cli_contract():
    """守卫契约：这些「宁可报错也不静默改写」的出口必须真可达，且给可执行提示。

    本 skill 的核心不变量是「绝不静默改写已有数据」，而守卫全部表现为
    `SystemExit` + 一句中文提示 —— **读代码测不到，只有跑真进程才守得住**。
    """
    import tempfile
    import shutil

    py = sys.executable
    bs = os.path.join(HERE, "build_matrix_rows.py")
    pr = os.path.join(HERE, "position_reuse.py")
    pb = os.path.join(HERE, "probe.py")

    d = tempfile.mkdtemp(prefix="wm_cli_")
    try:
        fx = _mk_offline_fixture(d)

        # 1) position_reuse：不说清「哪些行算本批新行」→ 必须报错退出（避免静默 0 条 / 误改历史行）
        rc, out, err = _run([py, pr, "--snapshot", fx["snap"],
                             "--out", os.path.join(d, "o.json"), "--config", fx["config"]])
        assert rc != 0, "position_reuse 未给 --new-rows/--start-row，却成功退出了"
        assert "未指定「本批新行」的范围" in (out + err), (out + err)[-300:]

        # 2) final：拿不到追加起始行 → 必须报错退出（不得兜底从第 2 行开始写）
        rc, out, err = _run([py, bs, "final", "--preview", fx["preview"],
                             "--config", fx["config"]])
        assert rc != 0, "final 拿不到追加起始行，却成功退出了"
        assert "无法确定追加起始行" in (out + err), (out + err)[-300:]

        # 3) 旧版 .xls：给可执行提示，**绝不抛裸 traceback**
        #    装了 openpyxl → 提示「另存为 .xlsx」；没装 → 先提示 pip install。两者都合格。
        xls = os.path.join(d, "old.xls")
        open(xls, "w", encoding="utf-8").close()
        rc, out, err = _run([py, pb, "workbook", "--workbook", xls])
        blob = out + err
        assert rc != 0, "读旧版 .xls 竟然成功退出了"
        assert "Traceback" not in blob, blob[-300:]
        assert ("另存为" in blob) or ("pip install openpyxl" in blob), blob[-300:]

        # 4) 各入口 --help 冒烟：argparse 接线没被改坏
        for argv in ([os.path.join(ROOT, "pipeline.py"), "--help"],
                     [bs, "final", "--help"],
                     [pr, "--help"],
                     [pb, "workbook", "--help"],
                     [os.path.join(HERE, "sheet_snapshot.py"), "build", "--help"],
                     [os.path.join(HERE, "to_kdocs_payload.py"), "--help"]):
            rc, out, err = _run([py] + argv)
            assert rc == 0, f"`{os.path.basename(argv[0])} {' '.join(argv[1:])}` 退出 {rc}：" \
                            f"{(err or out)[-200:]}"
    finally:
        shutil.rmtree(d, ignore_errors=True)


def smoke_offline_e2e():
    """离线端到端 golden 链路：preview CSV → final（快照通道）→ payload → kdocs rangeData。

    覆盖主链上「按列映射生成 payload / 算追加起始行 / 岗位复用 / 日期序列号 / 自动分批」
    这段**无真机就无法回归**的逻辑，全程不碰 openpyxl、网络与微信。
    """
    import json as _json
    import tempfile
    import shutil

    py = sys.executable
    bs = os.path.join(HERE, "build_matrix_rows.py")
    kd = os.path.join(HERE, "to_kdocs_payload.py")

    d = tempfile.mkdtemp(prefix="wm_e2e_")
    try:
        fx = _mk_offline_fixture(d)

        rc, out, err = _run([py, bs, "final", "--preview", fx["preview"],
                             "--snapshot", fx["snap"], "--out-dir", fx["out"],
                             "--config", fx["config"]])
        assert rc == 0, f"final 退出 {rc}：{(err or out)[-400:]}"
        # 追加起始行必须由快照算出（表头 1 + 历史 1 行 → 3），**不许兜底**
        assert "自动计算 = 3" in out, out[-400:]
        # 岗位复用：本批两行都是张三且岗位为空 → 从历史区补「客服」
        assert "补 2 个岗位" in out, out[-400:]

        vals = _json.load(open(os.path.join(fx["out"], "payload.json"),
                               encoding="utf-8"))["values"]
        cell = {(v["row"], v["col"]): v for v in vals}
        # 列序 = fixture 的 column_mapping：A=0 项目编号 B=1 需求描述 C=2 提出时间
        #                              D=3 提出人岗位 E=4 提出人 F=5 解决人
        assert cell[(2, 1)]["string_value"] == "把结算单的费用改成 1200", cell.get((2, 1))
        assert cell[(3, 1)]["string_value"] == "搜索箱号要等 15 秒", cell.get((3, 1))
        assert cell[(2, 3)]["string_value"] == "客服", cell.get((2, 3))
        assert cell[(3, 3)]["string_value"] == "客服", cell.get((3, 3))
        # 提出时间必须是 Excel 序列号（数字型），不是原始文本 2026-09-05
        assert cell[(2, 2)]["value_type"] == "NUMBER", cell.get((2, 2))
        assert cell[(2, 2)]["number_value"] == 46270, cell.get((2, 2))
        # 本批两行都必须落在第 3~4 行（不改写历史区）
        assert sorted({v["row"] for v in vals}) == [2, 3], sorted({v["row"] for v in vals})

        rc, out, err = _run([py, kd, "--payload", os.path.join(fx["out"], "payload.json"),
                             "--out", os.path.join(fx["out"], "kdocs_update.json"),
                             "--file-id", "F1", "--worksheet-id", "7",
                             "--config", fx["config"]])
        assert rc == 0, f"to_kdocs_payload 退出 {rc}：{(err or out)[-400:]}"

        upd = _json.load(open(os.path.join(fx["out"], "kdocs_update.json"), encoding="utf-8"))
        # 关键契约：工作表键名必须是连接器的 worksheet_id，不是 sheetId；
        # 且必须是**整数** —— CLI（字符串）与 config（整数）两条来源曾产出不同类型。
        assert "sheetId" not in upd, list(upd)
        assert upd["worksheet_id"] == 7 and isinstance(upd["worksheet_id"], int), \
            repr(upd["worksheet_id"])
        assert upd["call_count"] == len(upd["calls"]) and upd["call_count"] >= 1, upd["call_count"]
        for c in upd["calls"]:
            assert len(c["rangeData"]) <= upd["batch_size"], len(c["rangeData"])
        ops = [o for c in upd["calls"] for o in c["rangeData"]]
        assert len(ops) == upd["total_ops"], (len(ops), upd["total_ops"])

        formulas = [o for o in ops if o["opType"] == "formula"]
        assert len(formulas) == len(vals), (len(formulas), len(vals))
        # 每条 formula op 只能写一个值 → 必然单格（无法靠合并收敛）
        for o in formulas:
            assert o["rowFrom"] == o["rowTo"] and o["colFrom"] == o["colTo"], o
        assert {"把结算单的费用改成 1200", "客服"} <= {o["formula"] for o in formulas}
        # 日期列自动补 format op
        formats = [o for o in ops if o["opType"] == "format"]
        assert formats and all("numfmt" in o["xf"] for o in formats), formats
    finally:
        shutil.rmtree(d, ignore_errors=True)


def smoke_repo_hygiene():
    """仓库卫生：运行产物与本地绝对路径不得入库；全部跟踪文件必须是合法 UTF-8。

    这几项都是**真跑出来的**产物（探测/解密/导出/回填的中间物含聊天明文与账号目录），
    而 `.gitignore` 写错**不会报任何错** —— 只能靠断言守。git 不可用时跳过并说明。
    """
    import re as _re
    import fnmatch as _fnmatch

    try:
        pr = subprocess.run(["git", "-C", ROOT, "ls-files"], capture_output=True, timeout=90)
    except Exception as e:                          # 本机没装 git / git 不在 PATH
        raise Skip(f"没有可用的 git（{type(e).__name__}）")
    if pr.returncode != 0:
        raise Skip("当前目录不是 git 仓库")

    files = [f for f in pr.stdout.decode("utf-8", "replace").splitlines() if f.strip()]
    assert files, "git ls-files 返回空 —— 仓库里一个跟踪文件都没有？"

    # 1) 运行产物**一份名单两处用**：(a) 不得被跟踪；(b) `.gitignore` 必须逐条收录。
    #
    #    只做 (a) 是**结构性不够**的：`git ls-files` 只看**已跟踪**文件，
    #    "未跟踪且未忽略"的产物它根本看不见 —— 而那正是最可能被 `git add -A` 顺手带走的
    #    状态。这个盲区是**真栽过**的：`rules_check.py plan` 的默认 `--out` 是
    #    `<cwd>/rules_check_out`，从仓库根目录跑一次就把**渲染后的完整提示词**
    #    （含我方标识=微信号）写进了仓库，`git status` 冒出一个 `?? rules_check_out/`，
    #    而当时的名单与 `.gitignore` 都没提到它。
    #    所以这里补成**双向**：名单里的每一项，`.gitignore` 里都得找得到。
    never = [
        ("config.json", "本地个人配置（含账号目录/路径）"),
        ("wechat_pilot/", "解密/导出产物（账号目录与聊天明文）"),
        ("*.db", "微信数据库副本"),
        ("all_keys*.json", "解密密钥"),
        ("keys_*.json", "解密密钥"),
        ("sheet_snapshot.json", "云文档快照（含真实表格内容与文档 id）"),
        ("raw_*.json", "kdocs 读表原始返回"),
        ("kdocs_update*.json", "回填请求体"),
        ("payload*.json", "回填 payload"),
        ("merged_preview.csv", "裁决用预览（含需求原文）"),
        ("final_rows.csv", "最终行"),
        ("rules_check_out/", "判定行为验证产物（含渲染后的真实标识）"),
        ("prompt_rendered.txt", "渲染后的完整提示词（含我方标识=微信号）"),
        ("rules_cases/", "盲评用例（含渲染后的对方关键字）"),
        ("validation_report.txt", "产出契约校验报告（含 agent 文件名与违约值）"),
    ]

    def _match(rel, pat):
        base = os.path.basename(rel)
        if pat.endswith("/"):
            return rel.startswith(pat)
        if "*" in pat:
            return _fnmatch.fnmatch(base, pat)
        return base == pat or rel == pat

    gi_path = os.path.join(ROOT, ".gitignore")
    ignore_txt = open(gi_path, encoding="utf-8").read() if os.path.isfile(gi_path) else ""
    assert ignore_txt, ".gitignore 不存在或为空"

    hygiene = []
    for pat, why in never:
        hit = sorted({f for f in files if _match(f, pat)})
        if hit:
            hygiene.append(f"已被跟踪的运行产物 {pat}（{why}）：{hit[:5]}")
        if pat not in ignore_txt:
            hygiene.append(f"{pat} 没写进 .gitignore（{why}）—— 未跟踪且未忽略的产物"
                           "不会被 git ls-files 看见，`git add -A` 会直接带走")
    assert not hygiene, hygiene

    # 2) 本地绝对路径不得入库（用户名 / 账号目录名会随路径一起泄漏）
    #    用户名段限定为「字母/数字/._-」，这样正则源码与文档里的伪代码不会被误判成路径。
    abs_re = _re.compile(r"[A-Za-z]:[\\/]{1,2}Users[\\/][A-Za-z0-9._-]+"
                         r"|/Users/[A-Za-z0-9._-]+"
                         r"|/home/[A-Za-z0-9._-]+")
    # 扫描器自检：规则放宽（漏报真路径）与规则被写坏（误报）都是**静默**的，先钉住两头。
    # 样本用拼接而非字面量 —— 字面量会被自己扫到。
    assert abs_re.search("C:" + os.sep + "Users" + os.sep + "someone"), "漏报：真路径没被命中"
    assert not abs_re.search(r"/Users/[A-Za-z0-9._-]+"), "误报：正则源码被当成路径"

    hits = []
    for f in files:
        p = os.path.join(ROOT, f)
        if not os.path.isfile(p):
            continue
        try:
            txt = open(p, encoding="utf-8").read()     # 3) 顺带守「必须是合法 UTF-8」
        except UnicodeDecodeError as e:
            hits.append(f"{f}: 不是合法 UTF-8（{e.reason}）")
            continue
        m = abs_re.search(txt)
        if m:
            hits.append(f"{f}: 命中绝对个人路径 {m.group(0)!r}")
    assert not hits, hits

    # 2b) 文档里不得出现**具体的**子代理产出文件名。
    #     真机形态是 `agent<字母>_<账号名>.json`，而账号名往往就是**人名拼音** ——
    #     写一个实例到文档里，等于把真人姓名写进仓库（本仓的合规声明是"不含任何真实姓名"）。
    #     这条是**真栽过**的：我在 CHANGELOG 里为了说明"校验报告会含带人名拼音的文件名"，
    #     顺手写了个真机实例。只扫文档：`smoke_test` 自己的合成夹具（`agentA_ok.json` 等）在
    #     `.py` 里，不属于"文档示例"。示例一律用占位符。
    concrete = _re.compile(r"agent[A-Za-z]*_[A-Za-z0-9]+\.json")
    leaked = []
    for f in files:
        if not f.endswith((".md", ".txt", ".json")):
            continue
        p = os.path.join(ROOT, f)
        if not os.path.isfile(p):
            continue
        for n, line in enumerate(open(p, encoding="utf-8").read().splitlines(), 1):
            m = concrete.search(line)
            if m:
                leaked.append(f"{f}:{n} 出现具体的子代理产出文件名 {m.group(0)!r}"
                              "（真机形态带账号名/人名拼音）→ 请改用占位符，如 `agent<字母>_<账号>.json`")
    # 扫描器自检：漏报（真机形态没命中）与误报（占位符/脚本名被命中）两头都要钉
    _probe_leak = "agent" + "Z_" + "somebody" + ".json"
    assert concrete.search(_probe_leak), "漏报：真机形态没被命中"
    assert not concrete.search("agent_N.txt"), "误报：占位符被当成具体名"
    assert not concrete.search("split_for_agents.py"), "误报：脚本名被当成产出文件"
    assert not leaked, leaked

    # 4) 「产物落进 git 仓库」的告警必须**真的认得出仓库**（上一段是"名字有没有被纳管"，
    #    这一条是兜底机制本身）。它不是文件属性而是一次判断 —— 判据写反/写死都会**静默失效**，
    #    而那正是 G1 的漏法（默认 --out 落在仓库里，谁都没说话）。
    import contextlib as _ctx
    import io as _io
    import tempfile as _tempfile
    import shutil as _shutil
    from common import warn_if_inside_git_repo

    _d = _tempfile.mkdtemp(prefix="wm_hyg_")
    try:
        _repo = os.path.join(_d, "r")
        _plain = os.path.join(_d, "plain")
        os.makedirs(_repo)
        os.makedirs(_plain)
        gp = subprocess.run(["git", "init", "-q", _repo], capture_output=True, timeout=90)
        if gp.returncode != 0:
            print("  （跳过：临时目录里 git init 失败，无法验证告警判据）")
        else:
            def _quiet(fn):
                buf = _io.StringIO()
                with _ctx.redirect_stdout(buf):
                    r = fn()
                return r, buf.getvalue()

            hit_in, txt_in = _quiet(
                lambda: warn_if_inside_git_repo(os.path.join(_repo, "rules_check_out")))
            hit_out, txt_out = _quiet(
                lambda: warn_if_inside_git_repo(os.path.join(_plain, "rules_check_out")))
            assert hit_in, "产物落在 git 仓库内却没被认出来（告警失效）"
            assert "不要入库" in txt_in, f"告警没说清该怎么办：{txt_in!r}"
            assert not hit_out, f"仓库外的目录被误判成在仓库内：{txt_out!r}"
    finally:
        _shutil.rmtree(_d, ignore_errors=True)

    # 3) 换行必须与本仓库声明的 .gitattributes（eol=lf）一致。
    #    这条守卫是**踩过才加的**：用 Python 文本模式（`open(..., "w")`）重写文件时，
    #    Windows 会把每个 \n 悄悄变成 \r\n —— 内容看着没变、git 也会在下次 add 时归一，
    #    所以只有逐字节读才能发现。混用（CRLF 与 LF 并存）与纯 CRLF 都算违约。
    ga = os.path.join(ROOT, ".gitattributes")
    if os.path.isfile(ga) and "eol=lf" in open(ga, encoding="utf-8").read():
        mixed, pure_crlf = [], []
        for f in files:
            p = os.path.join(ROOT, f)
            if not os.path.isfile(p):
                continue
            b = open(p, "rb").read()
            if b"\r\n" in b:
                (mixed if b"\n" in b.replace(b"\r\n", b"") else pure_crlf).append(f)
        assert not mixed, f"这些文件换行混用（CRLF 与 LF 并存）：{mixed}"
        assert not pure_crlf, (
            f"仓库声明 eol=lf，但这些文件是纯 CRLF：{pure_crlf}\n"
            "  常见成因：在 Windows 上用 Python 文本模式写文件（\\n 被转成 \\r\\n）。")


def smoke_cp1252_stdio():
    """非 UTF-8 控制台（cp1252）下输出中文不能崩 —— Windows runner 的真实条件。

    实测 GitHub 的 `windows-latest` runner 标准流编码是 **cp1252**，而本 skill 的输出
    **全是中文**：没修就是 `print` **第一行**就 `UnicodeEncodeError` 崩掉，而
    ubuntu / git-bash / 本机工作区都是 UTF-8 ⇒ **本地永远复现不了，只有 CI 的 windows 格子会红**。
    这里用子进程把条件造出来，本地也能跑到。

    CI 里**故意不设** `PYTHONUTF8` / `PYTHONIOENCODING` —— 设了就把这个缺陷盖住了。
    """
    import tempfile
    import shutil

    py = sys.executable
    d = tempfile.mkdtemp(prefix="wm_cp1252_")
    try:
        fx = _mk_offline_fixture(d)
        env = dict(os.environ)
        env["PYTHONIOENCODING"] = "cp1252"
        cases = [
            # ① --help：argparse 的说明文字全中文
            ([os.path.join(ROOT, "pipeline.py"), "--help"], 0),
            ([os.path.join(HERE, "probe.py"), "workbook", "--help"], 0),
            # ② 真的会打中文错误：不给「本批新行」→ 多行中文提示 + exit 1
            ([os.path.join(HERE, "position_reuse.py"), "--snapshot", fx["snap"],
              "--out", os.path.join(d, "o.json"), "--config", fx["config"]], 1),
        ]
        for argv, want in cases:
            rc, out, err = _run([py] + argv, env=env)
            blob = out + err
            assert rc == want, (f"cp1252 下 `{os.path.basename(argv[0])}` 退出 {rc}"
                                f"（期望 {want}）：{blob[-300:]}")
            assert "UnicodeEncodeError" not in blob, blob[-300:]
    finally:
        shutil.rmtree(d, ignore_errors=True)


def smoke_local_channel_full():
    """--full：本地表格通道（openpyxl）读表 + 追加起始行 —— 用临时 xlsx，不碰真实文件。

    云文档通道的快照路径已被 `smoke_offline_e2e` 覆盖，本地通道此前**没有端到端覆盖**。
    """
    import json as _json
    import tempfile
    import shutil

    try:
        import openpyxl
    except ImportError:
        raise Skip("未安装 openpyxl —— 本地通道本就需要它")

    d = tempfile.mkdtemp(prefix="wm_xlsx_")
    try:
        p = os.path.join(d, "ledger.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "运维-2026"
        ws.append(["项目编号", "需求描述", "提出时间", "提出人岗位", "提出人"])
        ws.append(["P1", "历史甲", "2026-09-01", "客服", "张三"])
        ws.append(["P2", "历史乙", "2026-09-02", "", "张三"])
        wb.save(p)

        rc, out, err = _run([sys.executable, os.path.join(HERE, "probe.py"),
                             "workbook", "--workbook", p, "--json"])
        assert rc == 0, f"probe 退出 {rc}：{(err or out)[-300:]}"
        res = _json.loads(out)          # --json 模式下 stdout 必须是纯 JSON
        assert res["header_row"] == 1, res["header_row"]
        assert res["column_mapping"]["需求描述"] == "B", res["column_mapping"]
        assert res["column_mapping"]["提出人"] == "E", res["column_mapping"]
        assert res["date_columns"] == ["C"], res["date_columns"]
        assert res["last_data_row"] == 3, res["last_data_row"]
        assert res["next_append_row"] == 4, res["next_append_row"]
    finally:
        shutil.rmtree(d, ignore_errors=True)


def smoke_zstd_decode_full():
    """--full：ZSTD 压缩正文的**成功**解码路径。

    这是本 skill 代价最高的一条路径（实测 8 条乱码消息里 4 条含真实需求正文）。
    非 --full 只测得到"解不开时给明确标记"这个降级分支 —— **成功分支必须有库才跑得到**，
    所以装 zstandard 的 CI 格子不是陪跑，它守的就是这里。
    """
    try:
        import zstandard as _zstd
    except ImportError:
        raise Skip("未安装 zstandard —— 这是它唯一能加覆盖的地方")

    from export_conversations import render_content, decode_payload, ZSTD_MAGIC

    text = "把结算单的费用改成 1200（压缩正文示例）"
    blob = _zstd.ZstdCompressor().compress(text.encode("utf-8"))
    assert blob[:4] == ZSTD_MAGIC, blob[:4]                 # 构造的确实是 ZSTD 帧
    assert decode_payload(blob, 4) == (text, ""), decode_payload(blob, 4)
    assert render_content(blob, 4, 1) == text, render_content(blob, 4, 1)
    # 误判防护：即使 ct 没标 4，魔数也能识别出它是压缩内容
    assert decode_payload(blob, 0) == (text, ""), decode_payload(blob, 0)

    # 压缩的 appmsg（引用回复）：正文在 <refermsg> 里，必须整条还原而不是只剩 `[链接/文件]`
    xml = ('<?xml version="1.0"?><msg><appmsg><title>话头示例A</title>'
           '<refermsg><content>引用正文示例B</content></refermsg></appmsg></msg>')
    got = render_content(_zstd.ZstdCompressor().compress(xml.encode("utf-8")), 4, 49)
    assert got.startswith("[链接/文件] "), got
    assert "话头示例A" in got and "引用正文示例B" in got, got


def smoke_path_guards():
    """输入路径守卫：所有"文件不存在 / 坏 JSON"都必须**友好报错**，不许裸 traceback。

    背景：`open()` / `load_workbook()` 抛的是 FileNotFoundError，而项目对同类输入
    （`--snapshot`、`--raw`、`--payload`）一向给可执行提示 —— **同一类输入两种行为**本身就是缺陷。
    现全部收敛到 `common.require_file / load_json_file / open_local_workbook`，这里逐条钉住。
    """
    import tempfile
    import shutil

    py = sys.executable
    d = tempfile.mkdtemp(prefix="wm_paths_")
    try:
        fx = _mk_offline_fixture(d)
        nope = os.path.join(d, "NOPE.json")
        nope_xlsx = os.path.join(d, "NOPE.xlsx")
        bad = os.path.join(d, "bad.json")
        with open(bad, "w", encoding="utf-8") as f:
            f.write("{not json")

        cases = [
            ("probe --workbook", [os.path.join(HERE, "probe.py"), "workbook", "--workbook", nope_xlsx]),
            ("probe --snapshot", [os.path.join(HERE, "probe.py"), "workbook", "--snapshot", nope]),
            ("preview --src", [os.path.join(HERE, "build_matrix_rows.py"), "preview", "--src",
                               os.path.join(d, "NOPE_DIR"), "--out", os.path.join(d, "pv.csv"),
                               "--config", fx["config"]]),
            ("final --preview", [os.path.join(HERE, "build_matrix_rows.py"), "final", "--preview",
                                 os.path.join(d, "NOPE.csv"), "--snapshot", fx["snap"],
                                 "--out-dir", os.path.join(d, "o1"), "--config", fx["config"]]),
            ("position_reuse --snapshot", [os.path.join(HERE, "position_reuse.py"), "--snapshot", nope,
                                           "--start-row", "5", "--out", os.path.join(d, "o.json"),
                                           "--config", fx["config"]]),
            ("position_reuse --workbook", [os.path.join(HERE, "position_reuse.py"), "--workbook", nope_xlsx,
                                           "--start-row", "5", "--out", os.path.join(d, "o.json"),
                                           "--config", fx["config"]]),
            ("position_reuse --history", [os.path.join(HERE, "position_reuse.py"), "--history", nope,
                                          "--start-row", "5", "--out", os.path.join(d, "o.json"),
                                          "--config", fx["config"]]),
            ("position_reuse --new-rows", [os.path.join(HERE, "position_reuse.py"), "--snapshot", fx["snap"],
                                           "--new-rows", os.path.join(d, "NOPE.csv"),
                                           "--out", os.path.join(d, "o.json"), "--config", fx["config"]]),
            ("sheet_snapshot inspect", [os.path.join(HERE, "sheet_snapshot.py"), "inspect", "--snapshot", nope]),
            ("sheet_snapshot build --raw 坏JSON", [os.path.join(HERE, "sheet_snapshot.py"), "build",
                                                   "--raw", bad, "--out", os.path.join(d, "s.json")]),
            ("to_kdocs_payload --payload 坏JSON", [os.path.join(HERE, "to_kdocs_payload.py"), "--payload", bad,
                                                   "--out", os.path.join(d, "k.json"),
                                                   "--file-id", "F1", "--worksheet-id", "7"]),
            ("split_for_agents --transcripts", [os.path.join(HERE, "split_for_agents.py"),
                                                "--transcripts", os.path.join(d, "NOPE_DIR"), "--n", "2"]),
            # 坏 JSON（不是"不存在"）：配置文件 / 历史岗位文件写坏时也不许抛 JSONDecodeError
            ("rules_check plan --config 坏JSON", [os.path.join(HERE, "rules_check.py"), "plan",
                                                  "--config", bad]),
            ("position_reuse --history 坏JSON", [os.path.join(HERE, "position_reuse.py"),
                                                 "--history", bad, "--start-row", "5",
                                                 "--out", os.path.join(d, "o2.json"),
                                                 "--config", fx["config"]]),
            ("final --history 坏JSON", [os.path.join(HERE, "build_matrix_rows.py"), "final",
                                        "--preview", fx["preview"], "--snapshot", fx["snap"],
                                        "--history", bad, "--out-dir", os.path.join(d, "o3"),
                                        "--config", fx["config"]]),
            ("rules_check verify --results 缺失", [os.path.join(HERE, "rules_check.py"), "verify",
                                                   "--results", os.path.join(d, "NOPE_DIR"),
                                                   "--config", fx["config"]]),
        ]
        bad_cases = []
        for name, argv in cases:
            rc, out, err = _run([py] + argv, cwd=HERE)
            blob = out + err
            if rc == 0:
                bad_cases.append(f"{name}: 退出码 0（应报错）")
            elif "Traceback (most recent call last)" in blob:
                bad_cases.append(f"{name}: 抛了裸 traceback")
            elif not blob.strip():
                bad_cases.append(f"{name}: 没给任何提示")
        assert not bad_cases, bad_cases
    finally:
        shutil.rmtree(d, ignore_errors=True)


def smoke_no_silent_fallback():
    """**P0 回归守卫**：表格来源给了却不存在，必须报错退出，且**不得产出 payload**。

    曾把"路径打错"当成"没给来源"：`os.path.isfile()` 两道判断都不进 → `sheet=None` →
    起点静默改用 `config.excel.start_row` → 以 exit 0 产出 payload（回填到错误行）。
    这里把 config 的 start_row 设成 99，只要它还敢产出 payload 就说明退化回来了。
    """
    import json as _json
    import tempfile
    import shutil

    py = sys.executable
    d = tempfile.mkdtemp(prefix="wm_nofb_")
    try:
        fx = _mk_offline_fixture(d)
        cfg = _json.load(open(fx["config"], encoding="utf-8"))
        cfg["excel"]["start_row"] = 99                     # 诱饵：谁退到这里就会产出 payload
        with open(fx["config"], "w", encoding="utf-8") as f:
            _json.dump(cfg, f, ensure_ascii=False)

        for flag, missing in (("--workbook", os.path.join(d, "NOPE.xlsx")),
                              ("--snapshot", os.path.join(d, "NOPE.json"))):
            od = os.path.join(d, "out_" + flag.strip("-"))
            rc, out, err = _run([py, os.path.join(HERE, "build_matrix_rows.py"), "final",
                                 "--preview", fx["preview"], flag, missing,
                                 "--out-dir", od, "--config", fx["config"]])
            produced = os.path.isfile(os.path.join(od, "payload.json"))
            assert rc != 0, f"final {flag} <不存在> 竟然 exit 0：{(out + err)[-300:]}"
            assert not produced, f"final {flag} <不存在> 仍然产出了 payload.json（静默降级！）"
            assert "不存在" in (out + err), (out + err)[-300:]
    finally:
        shutil.rmtree(d, ignore_errors=True)


def smoke_preview_contract():
    """preview / final / new-rows 的**列契约**：列名不对要早报，不许静默跑出 0 条。

    两种真实来源：① 旧版产物用的是带列字母的列名（真机 `merged_preview.csv` 里是
    `提出人O / 岗位N / 提出日期M / 需求描述L`）；② 台账模板没有「提出人」列时
    `final_rows.csv` 也就没有这一列 → `position_reuse --new-rows` 读到 0 个人名。
    """
    import json as _json
    import tempfile
    import shutil

    py = sys.executable
    d = tempfile.mkdtemp(prefix="wm_pv_")
    try:
        fx = _mk_offline_fixture(d)

        # ① 旧格式 preview（列名带列字母）→ 友好报错，不 KeyError
        old = os.path.join(d, "old_preview.csv")
        with open(old, "w", newline="", encoding="utf-8-sig") as f:
            f.write("序号,会话,提出人O,岗位N,提出日期M,需求描述L,归类Q\n1,A,张三,,2026-09-05,改费用,需求\n")
        rc, out, err = _run([py, os.path.join(HERE, "build_matrix_rows.py"), "final",
                             "--preview", old, "--snapshot", fx["snap"],
                             "--out-dir", os.path.join(d, "o"), "--config", fx["config"]])
        assert rc != 0, f"旧格式 preview 竟然通过了：{(out + err)[-300:]}"
        assert "缺少必需列" in (out + err), (out + err)[-300:]
        assert "Traceback" not in (out + err), (out + err)[-300:]

        # ② --new-rows 缺「提出人」列 → 友好报错（旧实现静默 0 条 + exit 0）
        nr = os.path.join(d, "no_person.csv")
        with open(nr, "w", newline="", encoding="utf-8-sig") as f:
            f.write("项目编号,需求描述\nP1,改费用\n")
        rc, out, err = _run([py, os.path.join(HERE, "position_reuse.py"), "--snapshot", fx["snap"],
                             "--new-rows", nr, "--out", os.path.join(d, "o.json"),
                             "--config", fx["config"]])
        assert rc != 0, f"缺列 CSV 竟然 exit 0：{(out + err)[-300:]}"
        assert "缺少必需列" in (out + err), (out + err)[-300:]

        # ③ 只有表头、0 数据行的 preview → 说清"没有数据行"，不要误报"缺少必需列"
        #    （旧实现会安静地写出空 payload 并 exit 0；更早的实现会把它说成列名不对）
        empty = os.path.join(d, "empty_preview.csv")
        with open(empty, "w", newline="", encoding="utf-8-sig") as f:
            f.write(",".join(["序号", "会话", "提出人", "提出人岗位", "提出时间", "需求描述",
                              "需求归类", "影响级别", "优先级", "结果", "计划时间",
                              "完成时间", "跨会话标记", "证据节选", "备注"]) + "\n")
        rc, out, err = _run([py, os.path.join(HERE, "build_matrix_rows.py"), "final",
                             "--preview", empty, "--snapshot", fx["snap"],
                             "--out-dir", os.path.join(d, "oe"), "--config", fx["config"]])
        blob = out + err
        assert rc != 0, f"空 preview 竟然 exit 0：{blob[-300:]}"
        assert "没有任何数据行" in blob, f"应说清是空表而不是列名问题：{blob[-300:]}"
        assert "缺少必需列" not in blob, f"误报成列名不对：{blob[-300:]}"
    finally:
        shutil.rmtree(d, ignore_errors=True)


def smoke_q_values():
    """「需求归类」取值域三方一致：提示词 ↔ 代码白名单 ↔ 线上台账实际取值。

    实测口径：2026-09-18 读线上台账「需求归类」列 196 格 → 数据处理 104 / bug 31 /
    答疑 27 / 优化 24 / 需求 9。**优化与需求并存**，而 `优化或需求` 从不在台账里出现
    （它是提示词旧措辞"需求或优化"诱导出的连写值，由 QMAP 归一掉）。
    这类不一致**不抛异常**（只是把错值写进台账），所以只能靠断言守。
    """
    import re as _re
    import json as _json
    from build_matrix_rows import Q_VALUES, QMAP, OUTCOME_LABEL

    prompt = open(os.path.join(ROOT, "references", "agent-prompt-zh.txt"),
                  encoding="utf-8").read()
    m = _re.search(r'"Q":"([^"]+)"', prompt)
    assert m, "提示词输出格式里找不到 Q 枚举"
    enums = m.group(1).split("|")
    assert len(enums) == len(set(enums)), f"Q 枚举有重复：{enums}"

    # ① 提示词的 Q 枚举必须都落在「白名单 ∪ 归一映射的键」里，否则写进台账的就是没见过的值
    allowed = set(Q_VALUES) | set(QMAP)
    unknown = [e for e in enums if e not in allowed]
    assert not unknown, f"提示词枚举了台账不存在的值 {unknown}（台账实际取值：{list(Q_VALUES)}）"
    # ② 反向：白名单里的值提示词要能表达出来（否则某类需求永远归不了类）
    unreachable = [v for v in Q_VALUES if v not in enums]
    assert not unreachable, f"台账取值 {unreachable} 在提示词 Q 枚举里无法表达"

    # ③ 规则文本不能再出现"或"式连写措辞（它正是 `优化或需求` 的来源）
    assert "需求或优化" not in prompt, "提示词规则文本又写成「需求或优化」了（会诱导连写值）"

    # ④ QMAP 的键必须是**台账里没有**的值，否则等于把合法值改掉
    for k, v in QMAP.items():
        assert k not in Q_VALUES, f"QMAP 的键 {k} 已经算台账合法取值，不该再归一"
        assert v in Q_VALUES, f"QMAP 把 {k} 归一到 {v}，而 {v} 不在台账取值域内"

    # ⑤ Q 的取值域应与 SKILL.md 的记录一致（改代码忘改文档 → 这里红）
    skill = open(os.path.join(ROOT, "SKILL.md"), encoding="utf-8").read()
    for v in Q_VALUES:
        assert v in skill, f"SKILL.md 未记录「需求归类」的取值 {v}"
    assert OUTCOME_LABEL, "OUTCOME_LABEL 不该为空"


def smoke_doc_cli_flags():
    """文档里的 CLI 形状不能与 argparse 打架。

    具体守一条**踩过的**：`pipeline.py` 的 `--config` 是**全局**选项，必须写在子命令之前；
    参数速查表一度把它列在 `pipeline.py probe/run` 的参数里，照抄会得到
    `error: unrecognized arguments`（argparse 不会把子命令后的未知选项回溯给上层解析器）。
    """
    import re as _re
    import subprocess as _sp

    skill = open(os.path.join(ROOT, "SKILL.md"), encoding="utf-8").read()

    # ① 参数速查表里，pipeline 的行不得把 --config 列进"子命令的参数"
    offenders = [l.strip() for l in skill.splitlines()
                 if l.lstrip().startswith("| `pipeline.py ")
                 and "--config" in l and "pipeline.py --config" not in l]
    assert not offenders, f"参数速查表把全局 --config 写进了子命令参数：{offenders}"

    # ② 真的跑一遍，确认 argparse 的形状与上面一致（放在后面被拒、放在前面可用）
    rc_bad, out_bad, err_bad = _run([sys.executable, os.path.join(ROOT, "pipeline.py"),
                                     "probe", "--config", os.path.join(ROOT, "config.example.json")])
    assert rc_bad != 0 and "unrecognized arguments" in (out_bad + err_bad), \
        "预期 `probe --config …` 被 argparse 拒绝，实际：%s" % (out_bad + err_bad)[-200:]
    rc_ok, out_ok, err_ok = _run([sys.executable, os.path.join(ROOT, "pipeline.py"), "--help"])
    assert rc_ok == 0 and "--config" in out_ok, "全局 --config 应出现在 pipeline.py --help 里"


def smoke_rules_fixtures():
    """判定夹具的**完备性**（CI 能做的部分；真实判定由独立子代理跑，见 rules_check.py）。

    夹具在 `references/rules-fixtures.json`，用来支撑"改判定规则必须做行为验证"这套流程。
    这里只校**确定性**的东西：覆盖是否完整、标注是否自洽、锚点是否真实存在。
    7 条都是"不抛异常、只是悄悄失效"的类型，所以只能靠断言守。
    """
    import json as _json
    import re as _re
    from build_matrix_rows import OUTCOME_LABEL

    path = os.path.join(ROOT, "references", "rules-fixtures.json")
    fx = _json.load(open(path, encoding="utf-8"))
    cases = fx.get("cases") or []
    assert cases, "夹具是空的"
    skill = open(os.path.join(ROOT, "SKILL.md"), encoding="utf-8").read()
    cfg = _json.load(open(os.path.join(ROOT, "config.example.json"), encoding="utf-8"))
    rules_cfg = cfg.get("rules") or {}
    switches = ("ignore_if_rejected", "ignore_if_no_reply",
                "ignore_vague_complaints", "data_change_default_done")
    ignored = {"rejected", "no_reply", "vague"}

    # ① id 唯一、转录非空、rules 名真实存在、两态期望都在
    ids = [c["id"] for c in cases]
    assert len(ids) == len(set(ids)), f"id 有重复：{ids}"
    bad = []
    for c in cases:
        cid = c.get("id", "?")
        if not c.get("transcript"):
            bad.append(f"{cid}: transcript 为空")
        if not c.get("why"):
            bad.append(f"{cid}: 缺 why（人工复核要看）")
        for sw in c.get("rules") or []:
            if sw not in rules_cfg:
                bad.append(f"{cid}: rules 里的 {sw} 不在 config.example.json 的 rules 里")
        for st in ("on", "off"):
            if st not in (c.get("expect") or {}):
                bad.append(f"{cid}: 缺 expect.{st}")
    assert not bad, bad

    # ② 每个开关都被 ≥1 条样本覆盖，且该样本 on/off 期望**不同**（否则等于没测这个开关）
    uncovered = []
    for sw in switches:
        hit = [c for c in cases if sw in (c.get("rules") or [])
               and c["expect"]["on"] != c["expect"]["off"]]
        if not hit:
            uncovered.append(sw)
    assert not uncovered, f"这些开关没有任何「两态期望不同」的样本：{uncovered}"

    # ③ anchor 必须在 SKILL.md 里真实存在（防悬空引用：改章节/措辞后夹具悄悄失准）
    dangling = [f"{c['id']} → {c['anchor']!r}" for c in cases if c.get("anchor") not in skill]
    assert not dangling, f"夹具锚点在 SKILL.md 里找不到：{dangling}"

    # ④ 标注自洽：record=False 只能是「被忽略类」；完成类/未确认类必须 record=True
    for c in cases:
        for st in ("on", "off"):
            e = c["expect"][st]
            o, rec = e.get("outcome"), e.get("record")
            tag = f"{c['id']}.{st}"
            if c.get("not_a_requirement"):
                assert o is None and rec is False, f"{tag}: 非需求类应当 outcome=null 且不记录"
                continue
            assert o in OUTCOME_LABEL, f"{tag}: outcome {o!r} 不在 OUTCOME_LABEL 里"
            if not rec:
                assert o in ignored, f"{tag}: 标了不记录，但 outcome={o!r} 不是被忽略类"
            if o in ("done", "default_done", "pending"):
                assert rec, f"{tag}: outcome={o!r} 属于「有结论的行」，不能标成不记录"

    # ⑤ outcome 枚举要全覆盖（否则某个值改坏了也没样本发现）
    seen = {c["expect"][st]["outcome"] for c in cases for st in ("on", "off")} - {None}
    missing = sorted(set(OUTCOME_LABEL) - seen)
    assert not missing, f"这些 outcome 值没有样本覆盖：{missing}"

    # ⑥ 脱敏：夹具是公开文件，不得出现真实标识
    blob = _json.dumps(fx, ensure_ascii=False)
    leaks = []
    for pat, what in ((r"[A-Za-z]:[\\/]{1,2}Users[\\/][A-Za-z0-9._-]+", "绝对个人路径"),
                      (r"1[3-9]\d{9}", "手机号"),
                      (r"(?<!\d)\d{11,}(?!\d)", "≥11 位纯数字串")):
        m = _re.search(pat, blob)
        if m:
            leaks.append(f"{what}: {m.group(0)!r}")
    assert not leaks, f"夹具里疑似有真实标识：{leaks}"

    # ⑦ 规则没定义清楚的地方必须**显式登记**（而不是让夹具假装有答案）
    open_ids = {c["id"] for c in cases if c.get("open")}
    declared = {q.get("case") for q in (fx.get("_open_questions") or [])}
    assert open_ids == declared, \
        f"open 标记与 _open_questions 不一致：open={sorted(open_ids)} / 声明={sorted(declared)}"


def smoke_agent_contract():
    """判定子代理产出契约（分支②）：坏值必须在 **preview 阶段**被拦下。

    为什么值得单列一条：子代理的错值**不会当场报错**，而是顺着 preview → final 一路写进台账。
    旧实现只做宽松兜底（`r.get(...)`）：字段拼错、枚举越界、日期写成"5月14日"全部静默通过。
    这里逐类造一遍，确认四类硬拦 + 契约外字段只告警 + **没有跳过开关**。
    """
    import json as _json
    import tempfile
    import shutil

    py = sys.executable
    bs = os.path.join(HERE, "build_matrix_rows.py")
    d = tempfile.mkdtemp(prefix="wm_ctr_")
    try:
        ok = [{"chat": "agentA", "O": "甲", "N": "客服", "ask_date": "2026-05-14",
               "L": "改个费用", "Q": "需求", "R": "中", "S": "", "outcome": "done",
               "v_date": "2026-05-14", "w_date": "2026-05-14", "evidence": "原话", "note": ""}]
        with open(os.path.join(d, "agentA_ok.json"), "w", encoding="utf-8") as f:
            _json.dump(ok, f, ensure_ascii=False)
        # 非子代理产物（dict）：必须被跳过而不是被当成"顶层不是数组"拦下
        with open(os.path.join(d, "payload.json"), "w", encoding="utf-8") as f:
            _json.dump({"values": []}, f)

        # ① 干净输入 → 放行，且产出 preview + 报告
        out_ok = os.path.join(d, "out_ok")
        os.makedirs(out_ok, exist_ok=True)
        rc, out, err = _run([py, bs, "preview", "--src", d, "--out",
                             os.path.join(out_ok, "p.csv")])
        assert rc == 0, f"干净输入竟被拦：{(out + err)[-300:]}"
        assert os.path.isfile(os.path.join(out_ok, "p.csv")), "preview 未产出"
        assert os.path.isfile(os.path.join(out_ok, "validation_report.txt")), "报告未落盘"

        # ② 七类违约各一处 → 必须拦下，且不产出 preview
        bad = [{"chat": "agentB", "O": "", "N": "客服", "ask_date": "2026/9/5", "L": "",
                "Q": "优化或需求呗", "R": "高中", "S": "9", "outcome": "finished",
                "v_date": "5月14日", "w_date": "", "evidence": "", "note": "",
                "askdata": "2026-05-14"}]
        with open(os.path.join(d, "agentB_bad.json"), "w", encoding="utf-8") as f:
            _json.dump(bad, f, ensure_ascii=False)
        out_bad = os.path.join(d, "out_bad")
        os.makedirs(out_bad, exist_ok=True)
        rc, out, err = _run([py, bs, "preview", "--src", d, "--out",
                             os.path.join(out_bad, "p.csv")])
        blob = out + err
        assert rc != 0, f"坏产出竟然 exit 0：{blob[-300:]}"
        assert "拒绝生成 preview" in blob, blob[-300:]
        assert not os.path.isfile(os.path.join(out_bad, "p.csv")), "被拦却仍产出 preview"
        report = os.path.join(out_bad, "validation_report.txt")
        assert os.path.isfile(report), "完整清单没落盘"
        rep = open(report, encoding="utf-8").read()
        for kind in ("缺必填 O", "缺必填 L", "Q 取值越界", "outcome 取值越界",
                     "R 取值越界", "S 取值越界", "v_date 无法解析"):
            assert kind in rep, f"报告缺 {kind}：{rep[:300]}"
        # 契约外字段名只告警（出现在报告里），但它不该自己单独构成拒绝条件
        assert "契约外的字段名" in rep, rep[:300]

        # ③ 边界：agent*.json 写成 dict → "该产出没产出"，必须拦（而不是像 payload.json 那样跳过）
        with open(os.path.join(d, "agentX_broken.json"), "w", encoding="utf-8") as f:
            _json.dump({"values": []}, f)
        rc, out, err = _run([py, bs, "preview", "--src", d, "--out",
                             os.path.join(out_bad, "p.csv")])
        assert rc != 0 and "顶层不是数组" in (out + err), (out + err)[-300:]

        # ④ 没有逃生门：--skip-validation 必须不被 argparse 接受
        rc, out, err = _run([py, bs, "preview", "--src", d, "--out",
                             os.path.join(out_bad, "p.csv"), "--skip-validation"])
        assert rc != 0, "竟然存在 --skip-validation 逃生门"
    finally:
        shutil.rmtree(d, ignore_errors=True)


def _mk_kdocs_raw(path, rows, cols, extra=None):
    """造 kdocs 风格 raw：在**四角**放值，使 build 记下的 coverage 极值恰为 (rows, cols)。"""
    import json as _json
    rd = []
    for r in sorted({0, rows}):
        for c in sorted({0, cols}):
            rd.append({"rowFrom": r, "rowTo": r, "colFrom": c, "colTo": c,
                       "cellText": "表头" if (r, c) == (0, 0) else f"v{r},{c}",
                       "numFormat": "G/通用格式"})
    rd += list(extra or [])
    with open(path, "w", encoding="utf-8") as f:
        _json.dump({"code": 0, "message": "成功",
                    "data": {"detail": {"rangeData": rd}, "result": "ok"}}, f, ensure_ascii=False)


def smoke_snapshot_coverage():
    """分支④：云文档通道的**表内岗位复用**及其两条可判定检查。

    这条分支的前提曾被写错（"云文档读历史太贵、一律手写 --history"）—— 错在把"回包大小"
    当成了"进上下文的大小"（大回包会被宿主落盘）。断言分两半：
      * 机制：build 写 coverage；final 用 coverage 判「行够不到历史区 ⇒ 报错」/「列没覆盖 ⇒ 告警」；
        旧快照无 coverage 时不得误报；复用关掉时不检查。
      * 事实：云文档通道**真能**表内复用（历史 张三=客服 ⇒ payload 写入「客服」）。
    另加一条**文档反向断言**：不许再出现"云文档不做表内复用"这类已作废的结论。
    """
    import json as _json
    import tempfile
    import shutil

    py = sys.executable
    ss = os.path.join(HERE, "sheet_snapshot.py")
    bs = os.path.join(HERE, "build_matrix_rows.py")

    d = tempfile.mkdtemp(prefix="wm_cov_")
    try:
        fx = _mk_offline_fixture(d)          # config + preview 复用同一套（列映射 A~F）
        cfgp, prev = fx["config"], fx["preview"]

        def build_snap(name, rows, cols, extra=None, sheet="运维"):
            raw = os.path.join(d, name + ".raw.json")
            snap = os.path.join(d, name + ".snap.json")
            _mk_kdocs_raw(raw, rows, cols, extra)
            rc, out, err = _run([py, ss, "build", "--raw", raw, "--out", snap,
                                 "--sheet", sheet, "--worksheet-id", "3"])
            assert rc == 0, f"build 失败：{(out + err)[-300:]}"
            return snap

        def run_final(snap, tag, extra=None):
            od = os.path.join(d, "o_" + tag)
            os.makedirs(od, exist_ok=True)
            rc, out, err = _run([py, bs, "final", "--preview", prev, "--snapshot", snap,
                                 "--out-dir", od, "--config", cfgp] + (extra or []))
            return rc, out + err, os.path.isfile(os.path.join(od, "payload.json"))

        # ① build 必须把 coverage / trimmed 写进快照（裁边前的原始极值）
        snap_full = build_snap("full", 195, 4)
        sh = _json.load(open(snap_full, encoding="utf-8"))["sheets"][-1]
        assert sh.get("coverage") == {"rowFrom": 0, "rowTo": 195, "colFrom": 0, "colTo": 4}, sh.get("coverage")
        assert sh.get("trimmed") and sh["trimmed"]["rows"] == 196, sh.get("trimmed")

        # ② 快照只有 3 行 + 起始行 197 → 确定够不到历史区 ⇒ 报错且不产出 payload
        snap_tiny = build_snap("tiny", 2, 5)
        rc, blob, produced = run_final(snap_tiny, "tiny", ["--start-row-excel", "197"])
        assert rc != 0 and not produced, f"行够不到历史区却放行：rc={rc}\n{blob[-300:]}"
        assert "不可能" in blob and "历史区" in blob, blob[-300:]

        # ③ 同一个小快照，但关掉复用 → 与数据区无关，不该拦
        rc, blob, produced = run_final(snap_tiny, "tiny_noreuse",
                                       ["--start-row-excel", "197", "--no-reuse-position"])
        assert rc == 0 and produced, f"关掉复用仍被拦：{blob[-300:]}"

        # ④ 行读全、但列只到 C（不含岗位列 D）→ 只告警，不阻断
        snap_narrow = build_snap("narrow", 195, 2)
        rc, blob, produced = run_final(snap_narrow, "narrow", ["--start-row-excel", "197"])
        assert rc == 0 and produced, f"列没覆盖不该阻断：{blob[-300:]}"
        assert "没覆盖这两列" in blob, blob[-300:]

        # ⑤ 旧快照（无 coverage）→ 两条检查都跳过，不得误报
        old = _json.load(open(snap_full, encoding="utf-8"))
        for s in old["sheets"]:
            s.pop("coverage", None)
        snap_old = os.path.join(d, "old.snap.json")
        with open(snap_old, "w", encoding="utf-8") as f:
            _json.dump(old, f, ensure_ascii=False)
        rc, blob, produced = run_final(snap_old, "old", ["--start-row-excel", "197"])
        assert rc == 0 and produced, f"旧快照被误拦：{blob[-300:]}"
        assert "快照只覆盖到" not in blob, blob[-300:]

        # ⑥ 事实：云文档通道表内复用**真的生效**（这是本节结论的立足点）
        snap_reuse = build_snap("reuse", 2, 4, extra=[
            {"rowFrom": 1, "rowTo": 1, "colFrom": 4, "colTo": 4, "cellText": "张三",
             "numFormat": "G/通用格式"},                      # E 列 = 提出人
            {"rowFrom": 1, "rowTo": 1, "colFrom": 3, "colTo": 3, "cellText": "客服",
             "numFormat": "G/通用格式"},                      # D 列 = 提出人岗位
        ])
        rc, blob, produced = run_final(snap_reuse, "reuse")
        assert rc == 0, f"表内复用路径失败：{blob[-300:]}"
        # fixture 的 preview 有两行、提出人都是张三且岗位为空 → 两行都该补
        assert "补 2 个岗位" in blob, blob[-300:]
        vals = _json.load(open(os.path.join(d, "o_reuse", "payload.json"), encoding="utf-8"))["values"]
        got = [v.get("string_value") for v in vals if v.get("col") == 3]   # D 列
        assert got == ["客服", "客服"], got

        # ⑦ 文档反向断言：已作废的结论不许作为**现行说法**复活
        #   刻意按行判、而不是全文子串：允许它作为"曾如此、已作废"的历史说明留痕
        #   （补一个正确的结论而不说明"原来错在哪"，下一个人还会照旧走老路）。
        stale = ("云文档不做表内复用", "云文档通道不做表内岗位复用", "云文档拿不到")
        for fname in ("SKILL.md", os.path.join("references", "workflow-notes.md")):
            txt = open(os.path.join(ROOT, fname), encoding="utf-8").read()
            for n, line in enumerate(txt.splitlines(), 1):
                if any(s in line for s in stale) and not any(
                        m in line for m in ("作废", "曾", "修正", "早期版本")):
                    raise AssertionError(
                        f"{fname}:{n} 把已作废结论当现行说法：{line.strip()[:120]}")
        # "表内复用可行"的事实基础：第 2 趟读的关键列必须含这两列
        import sheet_snapshot as _ss
        assert "提出人" in _ss.KEY_COLUMNS and "提出人岗位" in _ss.KEY_COLUMNS, _ss.KEY_COLUMNS
    finally:
        shutil.rmtree(d, ignore_errors=True)


def smoke_skill_frontmatter():
    """SKILL.md frontmatter 的长度额度（平台侧上限 1024）。

    **为什么需要**：`description` 是路由真正读的字段，每次往里补触发词 / 改描述都在吃这个额度，
    而**超限不会报任何错**（不会拒绝加载、也不会截断告警）—— 只能靠断言守。
    兄弟 skill（`webbuilder-xwl-patch`）就是**靠人工撞线**才补上这条守卫的：
    加新触发词时把 description 撑到 **1029**，是人工量出来的。

    度量口径与兄弟 skill 对齐（**偏保守**）：取 frontmatter 里 `description` 的**原始文本**
    （含折行与缩进）与整个 frontmatter 块，两者都必须 ≤ 1024（**字符**）。
    ⚠️ 平台到底按**字符**还是**字节**计，我们**没有权威依据** —— 兄弟 skill 用的是字符，
    这里沿用同一口径，但同时把**字节**数一并报出来（中文 1 字 ≈ 3 字节，两者差 2 倍以上），
    任一口径用掉八成额度就提示。真到临界时请以平台实际行为为准，别只信这条断言。
    """
    import re as _re

    text = open(os.path.join(ROOT, "SKILL.md"), encoding="utf-8").read()
    m = _re.match(r"^---\r?\n(.*?)\r?\n---[ \t]*\r?\n", text, _re.S)
    assert m, "SKILL.md 开头没有 frontmatter 块（形如 --- / key: value / ---）"
    fm = m.group(1)

    # description 取值可能写成块标量（>- / > / |- / |）或同行标量；两种都要认，
    # 认不出就报错而不是静默按 0 处理 —— 否则改成另一种写法等于把守卫绕过去。
    block = _re.search(r"^description:[ \t]*(?:[>|][+-]?)?[ \t]*\r?\n((?:[ \t]+.*\r?\n?)+)", fm, _re.M)
    if block:
        raw = block.group(1)
    else:
        same = _re.search(r"^description:[ \t]*(.+)$", fm, _re.M)
        assert same, "frontmatter 里找不到 description（写法特殊，请同步本断言的解析）"
        raw = same.group(1)

    assert raw.strip(), "SKILL.md 的 description 是空的"
    dlen = len(raw)                                        # 与兄弟 skill 同口径（字符）
    val = _re.sub(r"\s+", " ", raw).strip()                 # 实际会被解析出的值
    assert dlen <= 1024, (
        f"SKILL.md 的 description 原始文本 {dlen} 字符 / {len(raw.encode('utf-8'))} 字节，"
        f"超过平台 1024 上限（归一后 {len(val)} 字符）—— 超限不会报错，只会在路由侧失效")
    assert len(fm) <= 1024, (
        f"SKILL.md 的 frontmatter 块 {len(fm)} 字符 / {len(fm.encode('utf-8'))} 字节，"
        f"超过平台 1024 上限（description 占 {dlen} 字符）")
    # 额度用掉大半时留个可见提示（不失败，但值得知道）—— 两种口径都看
    for label, used in (("description 字符", dlen),
                        ("description 字节", len(raw.encode("utf-8"))),
                        ("frontmatter 字符", len(fm)),
                        ("frontmatter 字节", len(fm.encode("utf-8")))):
        if used > 1024 * 0.8:
            print(f"  （提示：{label} 已用 {used}/1024，余 {1024 - used}）")


def smoke_rules_check_positive():
    """`rules_check.py` 的**正向路径**：渲染 + 盲评 case 生成 + `verify` 的三类差异分类。

    为什么单列一条：这两个子命令的**产出**此前零回归覆盖 —— smoke 里 `rules_check` 只出现过
    两次，**都是错误路径**（坏 JSON config / `--results` 缺失）。而 `verify` 的
    「误记 / 漏记 / 结果不符 + 待裁决 + 未判定」分类正是本支的核心价值，此前只手工跑过一次。
    分支② 的产出契约校验反而有完整断言，两边**不对称**。

    **真实判定仍由独立子代理跑、不进 CI**（要模型、非确定性，塞进去只会 flaky）；
    这里只守**确定性**的那一半：渲染有没有漏注入、盲评 case 会不会泄漏答案、分类判据对不对。
    """
    import json as _json
    import tempfile
    import shutil
    import rules_check as _rc

    py = sys.executable
    rc_py = os.path.join(HERE, "rules_check.py")
    d = tempfile.mkdtemp(prefix="wm_rcp_")
    try:
        cfg = _json.load(open(os.path.join(ROOT, "config.example.json"), encoding="utf-8"))
        cfg["people"].update({"handler": "运维甲", "my_identifiers": ["MY_ID_9f21", "甲"],
                              "counterparty_keyword": "乙公司", "service_object": "示例客户"})
        cfg["excel"]["business_system_description"] = "示例业务系统"
        cfgp = os.path.join(d, "config.json")
        _json.dump(cfg, open(cfgp, "w", encoding="utf-8"), ensure_ascii=False)

        _fx, cases = _rc.load_fixtures()
        assert len(cases) >= 10, f"夹具太少，这条断言失去意义：{len(cases)}"

        # ---- ① plan：渲染 + 盲评 case ----
        out = os.path.join(d, "out")
        rc, o, e = _run([py, rc_py, "plan", "--config", cfgp, "--out", out])
        assert rc == 0, f"plan 失败：{(o + e)[-400:]}"
        pr = open(os.path.join(out, "prompt_rendered.txt"), encoding="utf-8").read()
        assert "{{" not in pr, "渲染后仍残留占位符 —— 等于发出了裸模板"
        for v in ("MY_ID_9f21", "乙公司", "示例业务系统", "运维甲"):
            assert v in pr, f"占位符没注入真实值：{v} 不在渲染结果里"
        assert "开关" in pr, "（开关渲染表）没被注入"

        cdir = os.path.join(out, "rules_cases")
        got = sorted(os.listdir(cdir))
        assert len(got) == len(cases), f"case 数 {len(got)} ≠ 夹具数 {len(cases)}"
        for c in cases:
            p = os.path.join(cdir, f"case_{c['id']}.txt")
            assert os.path.isfile(p), f"缺 case 文件：{c['id']}"
            body = open(p, encoding="utf-8").read()
            # 盲评的**地基**：case 文件里不能有任何"答案"痕迹 —— 否则子代理照抄、验证就失效
            for banned in ("期望", "应记录", "应忽略", "误记", "漏记", "标注"):
                assert banned not in body, f"{c['id']} 的 case 文件出现 {banned!r}（可能泄漏答案）"
            assert c["why"][:16] not in body, f"{c['id']} 的 case 文件含夹具 why 原文"
            # 说话人 token 必须已渲染成真实标识
            assert "[我方]" not in body and "[对方]" not in body, f"{c['id']} 的说话人 token 没渲染"
            assert "运维甲" in body, f"{c['id']} 的 case 里没有处理人（渲染没生效）"

        # ---- ② verify：全部按标注 → 必须通过 ----
        rules = cfg.get("rules", {}) or {}
        state = {c["id"]: c["expect"][_rc._expected_state(c, rules)] for c in cases}
        ok_dir = os.path.join(d, "res_ok")
        os.makedirs(ok_dir)
        for c in cases:
            exp = state[c["id"]]
            n = exp.get("rows", 1) if exp["record"] else 0
            rows = [{"chat": c["id"], "O": "甲", "N": "", "ask_date": "2026-06-01",
                     "L": "占位描述", "Q": "需求", "R": "中", "S": "", "outcome": exp["outcome"],
                     "v_date": "2026-06-01", "w_date": "2026-06-01", "evidence": "", "note": ""}
                    for _ in range(n)]
            _json.dump(rows, open(os.path.join(ok_dir, f"case_{c['id']}.json"), "w",
                                  encoding="utf-8"), ensure_ascii=False)
        rc, o, e = _run([py, rc_py, "verify", "--config", cfgp, "--results", ok_dir])
        assert rc == 0, f"全按标注却报失败：{(o + e)[-400:]}"
        assert "全部按标注判定" in (o + e), (o + e)[-300:]

        # ---- ③ verify：三类差异各造一条 + 一条未判定，必须**分类**报出 ----
        import build_matrix_rows as _bm

        # 四个角色必须是**四条不同的**夹具：否则后面的 elif 会把前面那条盖掉
        # （我第一版就栽在这 —— c_miss 与 c_rows 选到了同一条，"结果不符"永远不触发）。
        _used = set()

        def _first(pred):
            c = next(x for x in cases if not x.get("open") and x["id"] not in _used and pred(x))
            _used.add(c["id"])
            return c

        c_mis = _first(lambda c: not state[c["id"]]["record"])                       # 应忽略
        c_miss = _first(lambda c: state[c["id"]]["record"])                          # 应记录
        c_rows = _first(lambda c: state[c["id"]]["record"] and state[c["id"]].get("outcome"))
        miss_id = _first(lambda c: True)["id"]                                       # 刻意不写文件
        other = next(x for x in sorted(_bm.OUTCOME_LABEL)
                     if x != state[c_rows["id"]]["outcome"])

        bad_dir = os.path.join(d, "res_bad")
        os.makedirs(bad_dir)
        for c in cases:
            exp = state[c["id"]]
            n = exp.get("rows", 1) if exp["record"] else 0
            out_c = exp["outcome"]
            if c["id"] == c_mis["id"]:
                n = 1                                             # 应忽略却记了 → 误记
            elif c["id"] == c_miss["id"]:
                n = 0                                             # 应记录却 0 行 → 漏记
            elif c["id"] == c_rows["id"]:
                n, out_c = max(n, 1), other                       # 行数对但 outcome 不符
            if c["id"] == miss_id:
                continue                                          # 不写文件 → 未判定
            rows = [{"chat": c["id"], "O": "甲", "N": "", "ask_date": "2026-06-01",
                     "L": "占位描述", "Q": "需求", "R": "中", "S": "", "outcome": out_c,
                     "v_date": "2026-06-01", "w_date": "2026-06-01", "evidence": "", "note": ""}
                    for _ in range(n)]
            _json.dump(rows, open(os.path.join(bad_dir, f"case_{c['id']}.json"), "w",
                                  encoding="utf-8"), ensure_ascii=False)

        rc, o, e = _run([py, rc_py, "verify", "--config", cfgp, "--results", bad_dir])
        blob = o + e
        assert rc != 0, f"三类差异 + 未判定竟然通过：{blob[-400:]}"
        for bucket, cid in (("误记", c_mis["id"]), ("漏记", c_miss["id"]),
                            ("结果不符", c_rows["id"]), ("未判定", miss_id)):
            assert f"✗ {bucket}：1" in blob, f"{bucket} 没被分类报出（或数量不对）：{blob[-500:]}"
            assert cid in blob, f"{bucket} 里没提到具体是哪条：{cid}"
    finally:
        shutil.rmtree(d, ignore_errors=True)


def smoke_doc_layering():
    """文档分层：`workflow-notes.md` **不得逐字复述** `SKILL.md` 的规则。

    为什么需要：这两个文件曾在「云文档接口」「纯文本风险」「岗位复用」「判定规则」等**六个主题上
    各写一份**、措辞还不同 —— 双份真相必然漂移（同一件事改一处漏一处，谁都不知道哪份是对的）。
    现已收敛成"`SKILL.md` 写规则、`workflow-notes.md` 写依据"，但那只是**写在文档里的约定**：
    没有守卫，下一个人复制粘贴一句规则过去，就又是双份。

    判据刻意只查**逐字重复的长行**（归一空白与标点后完全相同、且 ≥30 字符）：
    近似重复（同一件事的规则面 vs 依据面）是**允许**的，机器判不准，硬拦会天天误报。
    """
    import re as _re

    def _norm(s):
        return _re.sub(r"[\s*`（）()：:「」『』、，。；！？!?/｜|—\-]+", "", s)

    skill_txt = open(os.path.join(ROOT, "SKILL.md"), encoding="utf-8").read()
    notes_path = os.path.join(ROOT, "references", "workflow-notes.md")
    notes_txt = open(notes_path, encoding="utf-8").read()

    # 约定本身必须还在（否则守卫就没了依据）
    assert "分工" in notes_txt and "SKILL.md" in notes_txt.split("\n")[6], \
        "workflow-notes.md 开头缺少「与 SKILL.md 的分工」约定段"

    skill_norm = {_norm(l) for l in skill_txt.splitlines() if len(l.strip()) >= 30}
    infence = False
    dup = []
    for n, line in enumerate(notes_txt.splitlines(), 1):
        if line.lstrip().startswith("```"):
            infence = not infence
            continue
        if infence:
            continue
        s = line.strip()
        if len(s) < 30 or s.startswith(("#", "|", ">")):
            continue
        if _norm(s) in skill_norm:
            dup.append(f"workflow-notes.md:{n} 逐字复述了 SKILL.md 的内容：{s[:70]!r}")

    # 扫描器自检
    _probe_line = "8" * 0 + "这是一条来自 SKILL.md 的规则原句，长度足够触发重复判定"
    assert _norm(_probe_line) in {_norm(
        x) for x in ["  这是一条来自  SKILL.md 的规则原句，长度足够触发重复判定！"]}, "漏报：归一后应判为重复"
    assert _norm("另一句完全不同的话，只是长度也超过了三十个字符而已") not in skill_norm, "误报风险"
    assert not dup, dup


# ========================= 文档结构类守卫 =========================
# 起因：2026-09-18 的三步走（结构分层 → 引用/歧义/冲突 → 全流程）里，这些检查都只是**一次性探针**。
# 探针跑完就再没人跑，下一个人改文档照样踩同一个坑 —— 而且**真踩了**：结构重排把
# `workflow-notes.md` 的一节改名，却没同步那句指向它的「见「…」」，悬空引用就这样进了仓库，
# 而**现有守卫一条都看不见它**（它既不是坏 JSON，也不是不一致的枚举）。
# 引用/结构类缺陷的共同点是：**没人报错，只有读者会点空**。所以必须固化成常驻断言。
#
# 判据一律**收窄**到"能机器判定、几乎不误报"的形态 —— 守卫一旦开始误报，就会被当成噪音、然后被无视：
#   · 文件引用只认反引号里的**单扩展名** token（`*.md/.txt/.json` 这种"扩展名罗列"直接跳过）；
#   · 章节引用只认 `见/详见/参见 + 「…」` 这种**明确的指针句式**，不做全文「」扫描；
#   · 参数引用按**逐行归属**判（该行的参数必须属于该行那个脚本），而不是"某处存在"；
#   · 重名校验按**同一父节**判（CHANGELOG 每个版本都有 `### Added`，父节不同就不算重名）。
DOC_FILES = ("SKILL.md", "README.md", "CHANGELOG.md",
             "references/workflow-notes.md", "references/agent-prompt-zh.txt")
# 规范类文档：**CHANGELOG 刻意不在内** —— 它是历史记录，本来就该引用已作废的说法
# 与那些不随 skill 分发的审查脚本，拿"现行文档"的判据去要求它只会天天误报。
DOC_NORMATIVE = ("SKILL.md", "README.md", "references/workflow-notes.md")

_DOC_SEARCH_DIRS = ("", "scripts", "references", "references/tools", ".github/workflows")

# 文档里可以提到、但**本来就不在仓库里**的路径 —— 按定义不算悬空引用。
# ⚠ 运行期产物那几条**刻意不在这里**，而是从 `.gitignore` **派生**（见 `_allowed_by_gitignore`）：
#   本项目已经在"同一事实抄两份"上栽过多次（README 抄 `.gitignore` 清单漏列 15 项），
#   再抄一份"什么不在仓库里"名单 = 埋一个必然漂移的副本。
#   这里只留 `.gitignore` 覆盖不到的两类。
DOC_ALLOW_REFS = frozenset({
    # 交付给用户的示例路径：由用户自选位置，仓库里既不该有、也没法预先忽略
    "history_positions.json", "probe.json", "agent_N.txt",
    # 仓库外：第三方解密工具
    "wcdb_key_tool_windows.py", "wechat-passphrase.json",
    # 仓库外：历轮审查的探针与负向测试（CHANGELOG 记录它们存在过，但它们不随 skill 分发）
    "probe_error_paths.py", "probe_drift.py", "probe_refs.py", "probe_b4_coverage.py",
    "negative_test_b24.py", "negative_test_g1.py", "negative_test_g3.py",
    "final_regression_b24.py", "simulate_ci_checkout.py", "n_reuse_plan.py",
})
_DOC_ALLOW_REFS_RE = (
    re.compile(r"^agent_\d+\.json$"),        # 子代理产出按批命名（agent1.json / agent2.json …）
    re.compile(r"^agent_\d+\.txt$"),         # 分包清单
)
_GITIGNORE_GLOBS = None


def _allowed_by_gitignore(name):
    """`name` 是否被 `.gitignore` 覆盖 —— 被覆盖就意味着"它本来就不该在仓库里"。

    只取 `.gitignore` 里**不带斜杠**的条目（`config.json` / `raw_*.json` / `payload*.json` …）：
    带斜杠的是**目录**（`wechat_pilot/`），而文档引用的是目录里的文件路径，形态对不上。
    """
    global _GITIGNORE_GLOBS
    if _GITIGNORE_GLOBS is None:
        globs = []
        for l in open(os.path.join(ROOT, ".gitignore"), encoding="utf-8").read().splitlines():
            l = l.strip()
            if l and not l.startswith("#") and "/" not in l:
                globs.append(l)
        _GITIGNORE_GLOBS = globs
    return any(fnmatch.fnmatch(name, g) for g in _GITIGNORE_GLOBS)


def _doc_lines(rel):
    return open(os.path.join(ROOT, rel), encoding="utf-8").read().splitlines()


def _visible_lines(text):
    """返回 [(行号, 行)]，**剔除围栏代码块**。

    代码块里是命令、注释、示例，不是文档结构：`# 1) 第 2 趟…` 会被标题正则误当标题，
    块里的编号也不该按"有序列表连续"要求。凡是要判结构的检查都必须先过这一层。
    """
    out, infence = [], False
    for n, l in enumerate(text.splitlines(), 1):
        if l.lstrip().startswith("```"):
            infence = not infence
            continue
        if not infence:
            out.append((n, l))
    return out


def _headings_of_text(text):
    """返回 [(行号, 级别, 标题, 父标题)]；父标题 = 向上最近的**更高级**标题（没有则 None）。

    记父节是为了让"重名"判得准：`### Added` 在每个版本下都有，父节（版本号）不同就不是重名。
    用"全文查重 + 给 CHANGELOG 开例外"来绕反而更糟 —— 例外会挡住 CHANGELOG 里真正的重名。
    """
    vis = _visible_lines(text)
    heads = []
    for i, (n, l) in enumerate(vis):
        m = re.match(r"^(#{1,6})\s+(.+?)\s*$", l)
        if not m:
            continue
        lvl, title = len(m.group(1)), m.group(2).strip()
        parent = None
        for j in range(i - 1, -1, -1):
            mm = re.match(r"^(#{1,6})\s+(.+?)\s*$", vis[j][1])
            if mm and len(mm.group(1)) < lvl:
                parent = mm.group(2).strip()
                break
        heads.append((n, lvl, title, parent))
    return heads


def _doc_headings(rel):
    return _headings_of_text(open(os.path.join(ROOT, rel), encoding="utf-8").read())


def smoke_doc_refs():
    """文档里的引用必须解析到真实目标：**文件** / **markdown 链接** / **章节指针** / **参数速查表归属**。

    这四类都是"改一处、漏一处，且不会报任何错"的缺陷形状。特别是章节指针：
    重命名一节之后，指向它的那句话会静默失效 —— 读者点不到，而 CI 全绿。
    """
    prob = []

    def _resolves(tok):
        base = os.path.basename(tok)
        if base in DOC_ALLOW_REFS or any(rx.match(base) for rx in _DOC_ALLOW_REFS_RE):
            return True
        if _allowed_by_gitignore(base):
            return True
        return any(os.path.exists(os.path.join(ROOT, d, tok)) for d in _DOC_SEARCH_DIRS)

    # ---- A) 反引号里的文件引用 ----
    pat_file = re.compile(r"`([A-Za-z0-9_./\-]+\.(?:md|py|json|yml|yaml|txt|html|ini|cfg))`")
    a_checked = 0
    for rel in DOC_FILES:
        for n, line in enumerate(_doc_lines(rel), 1):
            for m in pat_file.finditer(line):
                tok = m.group(1)
                if not re.match(r"^[A-Za-z0-9_\-]+\.[a-z]{2,5}$", os.path.basename(tok)):
                    continue                     # `.md/.txt/.json` 这类"扩展名罗列"不是路径
                if tok.startswith(("http", "C:", "/")) or "*" in tok:
                    continue
                a_checked += 1
                if not _resolves(tok):
                    prob.append(f"[文件] {rel}:{n} 引用了不存在的 {tok}")

    # ---- B) markdown 链接 ----
    for rel in DOC_FILES:
        for n, line in enumerate(_doc_lines(rel), 1):
            for m in re.finditer(r"\[[^\]]*\]\(([^)\s]+)\)", line):
                tgt = m.group(1).split("#")[0]
                if not tgt or tgt.startswith(("http", "mailto:")):
                    continue
                if not os.path.exists(os.path.join(ROOT, tgt)):
                    prob.append(f"[链接] {rel}:{n} 指向不存在的 {tgt}")

    # ---- C) 章节指针 ----
    titles = set()
    for r, dirs, fs in os.walk(ROOT):
        dirs[:] = [d for d in dirs if d not in (".git", "__pycache__")]
        for f in fs:
            if not f.endswith(".md"):
                continue
            relp = os.path.relpath(os.path.join(r, f), ROOT).replace(os.sep, "/")
            for _n, _lvl, t, _p in _doc_headings(relp):
                titles.add(re.sub(r"[（(].*?[)）]", "", t).strip().strip("`").replace("**", ""))

    def _sec_resolves(name):
        # **只认"标题包含引用名"这一个方向**。反方向（引用名包含某个短标题）曾让守卫形同虚设：
        # 只要仓库里存在一个标题 "x"，`「任意带 x 的长句子」` 都会被判为可解析 —— 负向测试里
        # 那条"假章节名竟然解析成功"的自检断言就是这么被抓出来的。
        return any(name in x for x in titles)

    # 只认「」『』《》：这几种引号在中文文档里都用来指"某一节"。
    # 不认 markdown 的 `[x]`，那会和 `[文字](链接)` 撞车，误报比漏报更糟。
    # 已知盲区（刻意不处理）：`见“判定规则”`（弯引号）抓不到 —— 抓法越宽越容易误伤正文引用。
    ptr = re.compile(r"(?:见|详见|参见)[^\n]{0,6}[「『《]([^」』》]{2,30})[」』》]")
    c_checked = 0
    for rel in DOC_FILES:
        for n, line in enumerate(_doc_lines(rel), 1):
            for m in ptr.finditer(line):
                name = m.group(1).strip()
                c_checked += 1
                if not _sec_resolves(name):
                    prob.append(f"[章节] {rel}:{n} 「{name}」解析不到任何标题"
                                f"（改章节名后没同步引用？）")

    # ---- D) 参数速查表：该行的参数必须属于该行那个脚本 ----
    own = {}
    for r, dirs, fs in os.walk(ROOT):
        dirs[:] = [d for d in dirs if d not in (".git", "__pycache__")]
        for f in fs:
            if f.endswith(".py"):
                own[f] = set(re.findall(
                    r'"(--[A-Za-z0-9\-]+)"', open(os.path.join(r, f), encoding="utf-8").read()))

    skill = open(os.path.join(ROOT, "SKILL.md"), encoding="utf-8").read()
    sec = skill.split("### 参数速查")[1].split("### 可配置项")[0]
    d_checked = 0
    for line in sec.splitlines():
        m = re.match(r"\|\s*`([a-z_]+\.py)`[^|]*\|(.*)$", line)
        if not m:
            continue
        script, body = m.group(1), m.group(2)
        if script not in own:
            prob.append(f"[参数] 参数速查表列了磁盘上不存在的脚本 {script}")
            continue
        for flag in sorted(set(re.findall(r"--[A-Za-z0-9\-]+", body))):
            d_checked += 1
            if flag not in own[script]:
                prob.append(f"[参数] 参数速查表把 {flag} 归给了 {script}，但该脚本没这个参数")

    # ---- 扫描器自检：漏报与误报两头都要钉 ----
    assert not os.path.exists(os.path.join(ROOT, "no_such_script_xyz.py")), "自检样本名撞上真实文件"
    assert not _resolves("no_such_script_xyz.py"), "漏报：不存在的文件被判为可解析"
    assert _resolves("smoke_test.py"), "误报：真实存在的 smoke_test.py 被判为悬空"
    assert not _sec_resolves("绝对不存在的章节名xyz"), "漏报：假章节名竟然解析成功"
    assert _sec_resolves("判定规则"), "误报：真实章节「判定规则」解析失败"
    # `.gitignore` 派生必须真的在工作（不能悄悄退化成"什么都放行"或"什么都没放行"）
    assert _allowed_by_gitignore("config.json"), "派生失效：.gitignore 里明明有 config.json"
    assert _allowed_by_gitignore("raw_whatever.json"), "派生失效：raw_*.json 应当匹配"
    assert not _allowed_by_gitignore("smoke_test.py"), "派生过度：真实脚本被当成可忽略产物"
    assert _resolves("payload.json") and _resolves("config.json"), \
        "误报：运行期产物应由 .gitignore 派生放行"
    assert a_checked >= 10 and c_checked >= 5 and d_checked >= 20, \
        f"扫描器覆盖面塌了（文件 {a_checked} / 章节 {c_checked} / 参数 {d_checked}）—— 判据可能已失效"

    assert not prob, prob


def smoke_doc_structure():
    """文档自身的结构完整性：标题层级 / 同父重名 / 有序列表连续 / 目录树与清单 vs 磁盘。

    判据都是"**结构**"而非"措辞"，所以能机器判定且几乎不误报。
    最值钱的一条是"清单 vs 磁盘"：新加一个脚本、忘了同步文档里的清单表 —— 不会有任何报错。
    """
    prob = []

    # ---- ① 标题层级不得跳级 ----
    for rel in DOC_FILES:
        if not rel.endswith(".md"):
            continue
        prev = 0
        for n, lvl, title, _p in _doc_headings(rel):
            if prev and lvl > prev + 1:
                prob.append(f"[层级] {rel}:{n} h{prev} → h{lvl} 跳级：{title[:40]}")
            prev = lvl

    # ---- ② 同一父节下不得重名 ----
    for rel in DOC_FILES:
        if not rel.endswith(".md"):
            continue
        seen = {}
        for n, lvl, title, parent in _doc_headings(rel):
            key = (parent, title)
            if key in seen:
                prob.append(f"[重名] {rel}:{n} 父节「{parent}」下重复标题「{title}」"
                            f"（首次出现在第 {seen[key]} 行）")
            seen[key] = n

    # ---- ③ 顶层有序列表必须从 1 连续编号 ----
    # （提示词模板曾被【开关渲染表】从中间截断成 1,2,<表>,3,4,5,6 —— 编号错位但没人报错）
    for rel in ("SKILL.md", "references/agent-prompt-zh.txt"):
        run = []

        def _flush(run, rel=rel):
            if len(run) < 2:
                return
            nums = [x[1] for x in run]
            if nums != list(range(1, len(nums) + 1)):
                prob.append(f"[编号] {rel}:{run[0][0]} 起的有序列表编号不连续：{nums}")

        for n, l in _visible_lines(open(os.path.join(ROOT, rel), encoding="utf-8").read()):
            m = re.match(r"^(\d+)\.\s+\S", l)
            if m:
                run.append((n, int(m.group(1))))
            elif l.strip() == "" or l.startswith(("  ", "\t")):
                continue          # 空行与缩进续行都不算"列表结束"
            else:
                _flush(run)
                run = []
        _flush(run)

    # ---- ④ README 目录树 / SKILL.md 脚本清单表 vs 磁盘 ----
    readme = open(os.path.join(ROOT, "README.md"), encoding="utf-8").read()
    listed = set(re.findall(r"[├└]──\s+([A-Za-z0-9_.\-]+)",
                            readme.split("## 目录结构")[1].split("```")[1]))
    allnames = set()
    for r, dirs, fs in os.walk(ROOT):
        dirs[:] = [d for d in dirs if d != "__pycache__"]
        allnames |= set(dirs) | set(fs)
    ghost = sorted(x for x in listed if x not in allnames)
    assert not ghost, f"README 目录树列了磁盘上不存在的条目：{ghost}"
    unlisted_root = [x for x in sorted(os.listdir(ROOT))
                     if x not in (".git", "README.md") and x not in listed]
    assert not unlisted_root, f"README 目录树漏列了根目录条目：{unlisted_root}（新增文件要同步）"
    for sub in ("scripts", "references"):
        real = sorted(f for f in os.listdir(os.path.join(ROOT, sub)) if not f.startswith("."))
        miss = [f for f in real if f not in listed]
        assert not miss, f"README 目录树漏列 {sub}/：{miss}"

    skill = open(os.path.join(ROOT, "SKILL.md"), encoding="utf-8").read()
    sec = skill.split("### 脚本清单")[1].split("### 参数速查")[0]
    tabled = set(re.findall(r"^\|\s*`([a-z_]+\.py)`", sec, re.M))
    on_disk = set(f for f in os.listdir(os.path.join(ROOT, "scripts")) if f.endswith(".py"))
    on_disk.add("pipeline.py")                                  # 编排入口在根目录
    assert not (tabled - on_disk), f"脚本清单表列了磁盘上没有的脚本：{sorted(tabled - on_disk)}"
    assert not (on_disk - tabled), f"scripts/ 下有脚本没进清单表：{sorted(on_disk - tabled)}"

    # ---- ⑤ 新增的 .md 必须纳入守卫（否则"新加的文档"等于没人管）----
    real_md = set()
    for r, dirs, fs in os.walk(ROOT):
        dirs[:] = [d for d in dirs if d not in (".git", "__pycache__")]
        for f in fs:
            if f.endswith(".md"):
                real_md.add(os.path.relpath(os.path.join(r, f), ROOT).replace(os.sep, "/"))
    guarded = {f for f in DOC_FILES if f.endswith(".md")}
    assert not (real_md - guarded), \
        f"这些 .md 没被文档守卫覆盖（新增文档要加进 DOC_FILES）：{sorted(real_md - guarded)}"
    assert not (guarded - real_md), f"DOC_FILES 列了不存在的文档：{sorted(guarded - real_md)}"

    # ---- 扫描器自检 ----
    _synth = "# T\n\n```\n# not a heading\n```\n\n## A\n\n### B\n\n### B\n\n#### C\n"
    got = [x[2] for x in _headings_of_text(_synth)]
    assert got == ["T", "A", "B", "B", "C"], f"标题扫描器判错（含围栏跳过）：{got}"
    assert _headings_of_text(_synth)[4][3] == "B", "父节回溯错了"
    assert len(_doc_headings("SKILL.md")) >= 20, "SKILL.md 的标题数不合常理 —— 扫描器可能已失效"

    assert not prob, prob


def smoke_doc_claims():
    """文档里"声称的事实"必须与仓库现状一致：术语声明在位 / 作废说法不当现行 / 关键计数一致。

    计数类最值得守：同一个数字写在两三处，改一处漏一处**不会有任何报错**，
    本项目已经因此漏过两次（自检项数、夹具条数）—— 所以把它变成断言，而不是靠人记得。
    """
    import json as _json

    skill = open(os.path.join(ROOT, "SKILL.md"), encoding="utf-8").read()
    readme = open(os.path.join(ROOT, "README.md"), encoding="utf-8").read()

    # ---- ① 术语等价声明必须在位 ----
    # `kdocs` / `WPS 云文档` / `金山文档` 是同一个东西的几种叫法，新读者最容易在这里打结。
    brief = skill.split("## 30 秒速览")[1].split("**主链五步**")[0]
    assert "术语等价" in brief, "SKILL.md「30 秒速览」的「术语等价」段被删了"
    for term in ("kdocs", "WPS 云文档", "金山文档", "local"):
        assert term in brief, f"「术语等价」段缺了 {term} 的说明"

    # ---- ② 已作废的说法不许作为现行说法出现（按行判，留痕豁免）----
    STALE = ("云文档不做表内复用", "云文档通道不做表内岗位复用", "快照不含数据区",
             "云文档拿不到", "一律手写 `--history`", "需求或优化")
    EXEMPT = ("作废", "不再", "早期版本", "曾经", "旧结论", "过时", "不存在")
    WINDOW = 60          # 豁免词必须落在旧说法两侧这个范围内才算"留痕"

    def _stale_hit(line):
        """这一行是否把已作废的说法**当成现行说法**在讲。

        豁免词必须**贴近**被引述的旧说法。早先的判据是"同一行里出现豁免词即可"，对抗性尝试
        一测就穿：行尾随便写个"不存在"就能把作废结论重新放回来。
        现在改成 ±60 字符窗口 —— 本仓现有 5 处留痕全都在窗口内（真机核实过）。
        """
        for bad in STALE:
            i = line.find(bad)
            if i < 0:
                continue
            lo, hi = max(0, i - WINDOW), min(len(line), i + len(bad) + WINDOW)
            if not any(k in line[lo:hi] for k in EXEMPT):
                return True
        return False

    hits = []
    for rel in DOC_NORMATIVE:
        for n, line in enumerate(_doc_lines(rel), 1):
            if _stale_hit(line):
                hits.append(f"{rel}:{n} 把已作废的说法当现行：{line.strip()[:80]}")
    assert not hits, hits

    # ---- ③ 关键计数：文档写的 == 实际 ----
    fx = _json.load(open(os.path.join(ROOT, "references", "rules-fixtures.json"),
                         encoding="utf-8"))
    n_cases = len(fx["cases"])
    n_open = sum(1 for c in fx["cases"] if c.get("open"))
    m = re.search(r"夹具在 `references/rules-fixtures\.json`（\*\*(\d+) 条\*\*）", skill)
    assert m, "SKILL.md 里找不到「夹具…共 N 条」那句话（改了措辞要同步本断言）"
    assert int(m.group(1)) == n_cases, f"SKILL.md 写夹具 {m.group(1)} 条，实际 {n_cases} 条"
    m2 = re.search(r"当前有 \*\*(\d+) 条\*\*", skill)
    assert m2, "SKILL.md 里找不到「当前有 N 条」（open 夹具）那句话"
    assert int(m2.group(1)) == n_open, f"SKILL.md 写 open {m2.group(1)} 条，实际 {n_open} 条"

    yml = open(os.path.join(ROOT, ".github", "workflows", "selftest.yml"),
               encoding="utf-8").read()
    mat = yml.split("matrix:")[1].split("steps:")[0]
    lst = re.findall(r"^\s+([A-Za-z\-]+):\s*\[([^\]]*)\]", mat, re.M)
    if len(lst) != 3:
        raise Skip(f"CI matrix 不是 3 个内联列表（找到 {len(lst)}）—— 形状变了，本断言解析不了")
    grid = 1
    for _k, v in lst:
        grid *= len([x for x in v.split(",") if x.strip()])
    for rel, txt in (("SKILL.md", skill), ("README.md", readme)):
        for mm in re.finditer(r"共 (\d+) 格", txt):
            assert int(mm.group(1)) == grid, \
                f"{rel} 写 CI「共 {mm.group(1)} 格」，实际 {grid} 格（改 yml 要同步文档）"

    # ---- ④ README 的「当前版本」必须与 CHANGELOG 的最新版本段一致 ----
    # 刻意**不**校验"该版本是否有对应 git tag"：发版时机由人决定，本仓库常态是"版本已定、tag 未打"，
    # 那种断言会长期假红 —— 而**长期假红的守卫比没有守卫更糟**（会被当成噪音，然后被无视）。
    chg = open(os.path.join(ROOT, "CHANGELOG.md"), encoding="utf-8").read()
    m_chg = re.search(r"^## \[([^\]]+)\]", chg, re.M)
    assert m_chg, "CHANGELOG 里找不到版本段标题（形如 `## [1.1.0] - YYYY-MM-DD`）"
    m_rd = re.search(r"当前版本 \[`v?([\d.]+)`\]", readme)
    assert m_rd, "README 里找不到「当前版本 [`vX.Y.Z`]」那句（改了措辞要同步本断言）"
    assert m_rd.group(1) == m_chg.group(1), (
        f"README 说当前版本 {m_rd.group(1)}，CHANGELOG 最新段却是 {m_chg.group(1)} —— "
        f"发版时这两处必须一起改")

    # ---- ⑤ 不许把「已打 tag 的版本」写死在正文里 ----
    # tag 由发布者按时机打，**文档无法自证**；写死在正文里必然在发布那一刻变错，
    # 而没有任何检查会提醒你（我们刻意**不**校验 tag 存在性 —— 那种断言会长期假红）。
    # 所以这条守的不是"列表对不对"，而是"**不许依赖 tag 事实**"。
    TAGS = re.compile(r"已打 tag 的[^\n]{0,30}v\d")
    # 但**引述**旧写法是必需的（否则连"改了什么"都写不了）—— 与上面的 STALE 同一套判据：
    # 豁免标记必须落在被引述内容的 ±60 字符内，不能靠行尾随便写个词就豁免。
    QUOTED = ("原写", "原本写", "原文", "曾写", "之前写", "旧写法", "改为", "改成", "修正为", "不再")

    def _tag_hit(line):
        for m in TAGS.finditer(line):
            lo, hi = max(0, m.start() - 60), min(len(line), m.end() + 60)
            if not any(k in line[lo:hi] for k in QUOTED):
                return True
        return False

    for rel in ("CHANGELOG.md", "README.md"):
        for n, line in enumerate(_doc_lines(rel), 1):
            assert not _tag_hit(line), (
                f"{rel}:{n} 把已打 tag 的版本写死在正文里（发布那一刻必然与事实脱节；"
                f"引述旧写法请让「原写 / 改为」这类词靠近它）：{line.strip()[:80]}")

    # ---- 扫描器自检：留痕豁免两头都要钉 ----
    assert _stale_hit("云文档不做表内复用，一律走快照。"), "漏报：现行旧说法没被判为违规"
    assert not _stale_hit("> 早期版本写的是「云文档不做表内复用」，那个结论已作废。"), \
        "误报：留痕（引述旧结论并标明作废）被当成现行说法"
    assert _stale_hit("云文档不做表内复用。" + "（此处省略说明）" * 8 + "本条不存在歧义"), \
        "漏报：豁免词离得老远仍然生效 —— 判据退化成了「同行出现即可」"
    assert _tag_hit("已打 tag 的是 `v1.0.0` 与 `v1.0.1`，请据此 checkout。"), \
        "漏报：写死的 tag 枚举没被抓到"
    assert not _tag_hit("**已打 tag 的版本以文末「版本链接」区为准，此处不列举**"), \
        "误报：不写死枚举的表述被当成违规"
    assert not _tag_hit("CHANGELOG 头部原写「已打 tag 的是 v1.0.0 与 v1.0.1」，发布同一分钟就错了"), \
        "误报：引述旧写法被当成了现行依赖"
    assert _tag_hit("已打 tag 的是 v1.0.0。" + "（补记若干说明）" * 8 + "以上请以实际为准。"), \
        "漏报：豁免词离得老远仍然生效"
    assert n_cases >= 10, f"夹具太少（{n_cases}），这几条计数断言失去意义"


def main():
    full = "--full" in sys.argv
    from common import ensure_utf8_stdio
    ensure_utf8_stdio()          # 必须在**任何输出之前**：否则 Windows 的 cp1252 下第一行就崩
    print("== import 冒烟 ==")
    check(f"{len(MODULES)} 个模块 import", smoke_imports)
    print("== common 单测 ==")
    check("serial/col_index/col_letter", smoke_common)
    print("== 消息正文解码（真机乱码根因）==")
    check("ZSTD 降级/appmsg 引用/sysmsg 归一化", smoke_message_decoding)
    print("== 守卫与稳健性（真跑暴露）==")
    check("空映射不给追加行/剥离不剥空/null 日期不崩", smoke_guards_and_robustness)
    print("== 云文档快照抽象层 ==")
    check("稀疏->密集网格 / GridSheet 接口 / 追加行", smoke_snapshot_grid)
    check("raw.json 多形态识别", smoke_snapshot_raw_parse)
    check("payload -> kdocs rangeData", smoke_kdocs_payload)
    check("probe.analyze_workbook 吃 GridWorkbook", smoke_probe_analyze)
    print("== 岗位复用共用层（final 内置 / position_reuse 共用）==")
    check("列解析/历史只读到 end_row/只填空缺/批内传播", smoke_reuse_position)
    print("== build_matrix_rows 单测 ==")
    check("norm_o/ids/QMAP", smoke_build)
    print("== 判定规则一致性（防漂移）==")
    check("outcome 标签 + 规则开关三处同步", smoke_rule_consistency)
    print("== 「需求归类」取值域（提示词 ↔ 代码 ↔ 线上台账）==")
    check("Q 枚举/白名单/归一映射/文档四处一致", smoke_q_values)
    print("== 判定夹具完备性（行为验证的 CI 侧；真实判定由 rules_check.py 跑）==")
    check("覆盖/自洽/锚点/枚举/脱敏/open 登记 共 7 项", smoke_rules_fixtures)
    print("== rules_check 的正向路径（渲染 / 盲评 case / verify 分类）==")
    check("渲染无残留占位符 + case 不泄答案 + 三类差异分类", smoke_rules_check_positive)
    print("== 文档里的 CLI 形状 vs argparse ==")
    check("全局 --config 的位置（写错位置会被 argparse 拒）", smoke_doc_cli_flags)
    print("== CLI 守卫契约（宁可报错也不静默改写）==")
    check("缺关键输入即报错 / .xls 给可执行提示 / 入口 --help", smoke_cli_contract)
    print("== 输入路径守卫（路径不存在/坏 JSON 一律友好报错）==")
    check("16 个场景：无裸 traceback、退出码非 0、有提示", smoke_path_guards)
    print("== 不许静默降级（P0 回归）==")
    check("表格来源不存在时不产出 payload", smoke_no_silent_fallback)
    print("== preview / final / --new-rows 的列契约 ==")
    check("旧格式或缺列的输入早报，不静默 0 条", smoke_preview_contract)
    print("== 子代理产出契约（坏值不许顺到 final）==")
    check("四类硬拦 + 契约外字段只告警 + 无跳过开关", smoke_agent_contract)
    print("== 云文档通道表内复用 + 快照覆盖范围 ==")
    check("coverage 元数据 / 两条检查 / 复用真生效 / 旧结论不复活", smoke_snapshot_coverage)
    print("== SKILL.md frontmatter 长度额度（平台上限 1024）==")
    check("description 与整块都不超限（超限不报错，只能靠断言守）", smoke_skill_frontmatter)
    print("== 文档分层（SKILL.md 写规则 / workflow-notes 写依据）==")
    check("notes 不得逐字复述 SKILL.md 的规则", smoke_doc_layering)
    print("== 文档引用可解析（文件 / 链接 / 章节指针 / 参数速查表归属）==")
    check("四类引用都指向真实目标（改章节名后漏同步会被抓住）", smoke_doc_refs)
    print("== 文档结构（层级 / 同父重名 / 编号连续 / 清单 vs 磁盘）==")
    check("结构与清单都在磁盘上兑现（新增文件必须同步文档）", smoke_doc_structure)
    print("== 文档声明 vs 事实（术语 / 作废说法 / 关键计数）==")
    check("术语声明在位、计数与夹具/CI 实际数一致", smoke_doc_claims)
    print("== 离线端到端（preview CSV → final → payload → kdocs rangeData）==")
    check("列映射/起始行/岗位复用/日期序列号/分批", smoke_offline_e2e)
    print("== 仓库卫生 ==")
    check("运行产物与绝对路径不入库 / 全文件合法 UTF-8", smoke_repo_hygiene)
    print("== 非 UTF-8 控制台（cp1252，Windows runner 的真实条件）==")
    check("输出中文不崩（含真打中文错误的守卫路径）", smoke_cp1252_stdio)
    print("== 依赖缺失时的提示 ==")
    check("缺 openpyxl 给可执行提示（非裸 ImportError）", smoke_openpyxl_hint)
    print("== position_reuse ==")
    check("import position_reuse", smoke_position_reuse)
    if full:
        check("positions/快照通道 Cell 读取", smoke_position_reuse_openpyxl)
        check("本地表格通道：probe 读 xlsx + 追加起始行", smoke_local_channel_full)
        check("ZSTD 压缩正文的成功解码（含压缩的引用回复）", smoke_zstd_decode_full)
    else:
        print("  （加 --full 启用附加检查）")

    # ---- 元检查：文档里声明的项数必须等于实际项数（不计入它自己）----
    # README 记的是**默认模式**（不带 --full）的项数，所以 --full 下跳过。
    _total = STAT["ok"] + STAT["skip"] + len(FAIL)

    def _doc_count_matches():
        if full:
            raise Skip("--full 的项数与 README 记的默认项数不同，用默认模式核")
        import re as _re
        txt = open(os.path.join(ROOT, "README.md"), encoding="utf-8").read()
        m = _re.search(r"smoke_test\.py\s+#\s*(\d+)\s*项", txt)
        assert m, "README 的「自检」节里没写「N 项」"
        assert int(m.group(1)) == _total, \
            f"README 写的是 {m.group(1)} 项，实际是 {_total} 项（改自检后要同步 README）"

    print("== 自检项数 vs 文档 ==")
    check("README 声明的项数 == 实际项数（防文档漂移）", _doc_count_matches, count=False)

    # 另一头：`--full` 相对默认**多跑几项**也要与 README 一致。只守默认数会漏掉
    # "加了一条 --full 专属检查、却忘了改文档里的增量" —— 两个方向各在一个模式下核才闭环。
    def _doc_full_delta_matches():
        if not full:
            raise Skip("「--full 再加 N 项」只能在 --full 下核（默认模式拿不到增量）")
        import re as _re
        txt = open(os.path.join(ROOT, "README.md"), encoding="utf-8").read()
        m = _re.search(r"smoke_test\.py\s+#\s*(\d+)\s*项", txt)
        m2 = _re.search(r"--full\s*#\s*再加\s*(\d+)\s*项", txt)
        assert m and m2, "README 的「自检」节里缺「N 项」或「--full 再加 N 项」"
        delta = _total - int(m.group(1))
        assert int(m2.group(1)) == delta, (
            f"README 说 --full 再加 {m2.group(1)} 项，实际加了 {delta} 项（改自检后要同步 README）")

    check("README 的「--full 再加 N 项」== 实际增量（另一头）", _doc_full_delta_matches, count=False)

    print()
    if FAIL:
        print(f"✗ 冒烟失败 {len(FAIL)} 项（通过 {STAT['ok']}，跳过 {STAT['skip']}）：{FAIL}")
        sys.exit(1)
    tail = f"，{STAT['skip']} 项跳过" if STAT["skip"] else ""
    print(f"✓ 全部通过（{STAT['ok']} 项{tail}）")
    if STAT["skip"]:
        print("  注意：有检查被跳过 —— CI 的 deps=full 格子就是为保证它们**都能跑**。")


if __name__ == "__main__":
    main()
