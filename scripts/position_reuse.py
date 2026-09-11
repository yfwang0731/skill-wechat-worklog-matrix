# -*- coding: utf-8 -*-
"""岗位列向上复用计划：按行序向上查同一"提出人"最近的"岗位"值，用于填补空缺。

**列位置由 config.json 的 column_mapping 决定**（提出人 / 提出人岗位），不再写死 N/O。
默认只填空缺（不覆盖已有值）；加 --override 才允许覆盖为上方最近值。
输出 payload（editor_sdk sheet_set_range_value 格式）供回填；
云文档通道再经 to_kdocs_payload.py 转成 kdocs rangeData。

依赖：openpyxl（仅本地通道；快照通道不需要）
用法：
  python position_reuse.py --workbook <xlsx> --out payload_n.json [--override]     # 本地
  python position_reuse.py --snapshot <json> --out payload_n.json [--override]     # 云文档
"""
import json
import sys
import argparse

from common import load_config, col_index


def main():
    ap = argparse.ArgumentParser()
    src = ap.add_mutually_exclusive_group(required=True)
    src.add_argument("--workbook", help="本地 .xlsx 路径（本地通道）")
    src.add_argument("--snapshot", help="云文档快照 json（WPS 通道）")
    ap.add_argument("--out", default="payload_n.json")
    ap.add_argument("--override", action="store_true")
    ap.add_argument("--start-row", type=int, default=2)
    ap.add_argument("--end-row", type=int, default=None)
    ap.add_argument("--config", default=None)
    args = ap.parse_args()

    cfg, _ = load_config(args.config)
    cfg = cfg or {}
    if (cfg.get("rules", {}) or {}).get("reuse_position_column") is False:
        print("[skip] config.rules.reuse_position_column=false，岗位列复用已跳过。")
        with open(args.out, "w", encoding="utf-8") as f:
            json.dump({"values": []}, f, ensure_ascii=False)
        sys.exit(0)
    excel_cfg = cfg.get("excel", {}) or {}
    mapping = excel_cfg.get("column_mapping") or {}
    sheet_hint = excel_cfg.get("sheet_match") or "运维"
    header_row = excel_cfg.get("header_row") or 1

    # 列定位：优先按表头映射，缺失时退回常见位置（O=提出人, N=岗位）
    if not mapping.get("提出人") or not mapping.get("提出人岗位"):
        print("⚠ 警告：config.excel.column_mapping 未包含「提出人/提出人岗位」，"
              "将退回默认列位（提出人=O, 岗位=N）。建议先跑 probe.py workbook 生成映射。")
    col_person = col_index(mapping["提出人"]) + 1 if mapping.get("提出人") else 15
    col_post = col_index(mapping["提出人岗位"]) + 1 if mapping.get("提出人岗位") else 14
    print(f"[columns] 提出人=第{col_person}列  岗位=第{col_post}列（来源："
          f"{'表头映射' if mapping.get('提出人') else '默认位置'}）")

    if args.snapshot:
        from common import load_snapshot, workbook_from_snapshot, pick_sheet
        wb = workbook_from_snapshot(load_snapshot(args.snapshot))
        sheet = pick_sheet(wb, sheet_hint)
        print(f"[source] WPS 快照 {args.snapshot}（worksheet_id={getattr(sheet, 'sheet_id', None)}）")
    else:
        import openpyxl
        from openpyxl.utils.exceptions import InvalidFileException
        # 旧版 .xls/.xlt 二进制格式 openpyxl 不支持（.xlsx/.xlsm/.xltx 可以）
        if args.workbook.lower().endswith((".xls", ".xlt")):
            sys.exit("✗ openpyxl 不支持旧版 .xls/.xlt 格式。请先用 Excel/WPS 把工作簿「另存为 .xlsx」再运行。")
        try:
            wb = openpyxl.load_workbook(args.workbook)
        except InvalidFileException:
            sys.exit(f"✗ 无法以 .xlsx 解析 {args.workbook}。若是旧版 .xls，请先另存为 .xlsx 再运行。")
        names = [s for s in wb.sheetnames if sheet_hint in s]
        sheet = wb[names[0] if names else wb.sheetnames[0]]
    end = args.end_row or sheet.max_row

    last_post = {}
    changed = []
    start = max(args.start_row, header_row + 1)
    for r in range(header_row + 1, start):      # 历史区：仅记录，不改
        person = sheet.cell(row=r, column=col_person).value
        post = sheet.cell(row=r, column=col_post).value
        if person and str(person).strip() and post and str(post).strip():
            last_post[str(person).strip()] = str(post).strip()

    for r in range(start, end + 1):
        person = sheet.cell(row=r, column=col_person).value
        post = sheet.cell(row=r, column=col_post).value
        if not person or not str(person).strip():
            continue
        person = str(person).strip()
        cur = str(post).strip() if post else ""
        if person in last_post and last_post[person]:
            hist = last_post[person]
            if (not cur) or (args.override and cur != hist):
                changed.append((r, person, cur, hist))
                cur = hist
        if cur:
            last_post[person] = cur

    cells = [{"row": r - 1, "col": col_post - 1, "value_type": "STRING", "string_value": new}
             for r, person, old, new in changed]
    with open(args.out, "w", encoding="utf-8") as f:
        json.dump({"values": cells}, f, ensure_ascii=False)
    print(f"sheet: {sheet.title} | changed: {len(changed)}")
    for c in changed:
        print(f"  R{c[0]} {c[1]}: '{c[2]}' -> '{c[3]}'")
    print("payload ->", args.out)


if __name__ == "__main__":
    main()
